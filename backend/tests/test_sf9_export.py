from datetime import date, datetime, timezone
import io
import uuid
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.v1.routes.Classes import router as classes_router
from app.core.Dependencies import get_db, get_staff_id
from app.db.Base import Base
from app.models.academic.AcademicLevel import AcademicLevel
from app.models.academic.AcademicPeriod import AcademicPeriod
from app.models.academic.AcademicYear import AcademicYear
from app.models.academic.Class_ import Class
from app.models.academic.StudentCLass import StudentClass
from app.models.academic.StudentPeriodGrade import StudentPeriodGrade
from app.models.academic.Subject import Subject
from app.models.academic.SubjectLoad import SubjectLoad
from app.models.attendance.Attendance import AttendanceRecord
from app.models.auth.UserAccount import UserAccount
from app.models.people.AcademicStaff import AcademicStaff
from app.models.people.Student import Student
from app.models.settings.Setting import Setting
from app.services.export.SF9ExportService import export_advisory_class_sf9, export_student_sf9
from fastapi import FastAPI

TABLES = [
    AcademicYear.__table__,
    AcademicLevel.__table__,
    AcademicPeriod.__table__,
    Class.__table__,
    Subject.__table__,
    SubjectLoad.__table__,
    UserAccount.__table__,
    AcademicStaff.__table__,
    Student.__table__,
    StudentClass.__table__,
    StudentPeriodGrade.__table__,
    AttendanceRecord.__table__,
    Setting.__table__,
]


def setup_test_db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine, tables=TABLES)
    Session = sessionmaker(bind=engine)
    db = Session()

    # School Settings
    db.add_all([
        Setting(key="school_name", value="Medellin National Science and Technology School (MNSTS)", type="string"),
        Setting(key="school_id", value="303012", type="string"),
        Setting(key="school_region", value="Region VII", type="string"),
        Setting(key="school_division", value="Cebu Province", type="string"),
        Setting(key="principal_name", value="Dr. Juanita Principal", type="string"),
    ])

    # Academic Structure
    ay = AcademicYear(academic_year_id=1, year_label="2026-2027", start_date=date(2026, 8, 1), end_date=date(2027, 5, 31), is_active=True)
    level = AcademicLevel(academic_level_id=1, level_name="Grade 8", grade_level=8)
    p1 = AcademicPeriod(academic_period_id=1, academic_year_id=1, period_name="Quarter 1", period_sequence=1, start_date=date(2026, 8, 1), end_date=date(2026, 10, 15), is_active=True)
    p2 = AcademicPeriod(academic_period_id=2, academic_year_id=1, period_name="Quarter 2", period_sequence=2, start_date=date(2026, 10, 16), end_date=date(2026, 12, 15), is_active=False)

    # Adviser & Non-adviser Teachers
    user1_id = uuid.uuid4()
    user2_id = uuid.uuid4()
    adviser_user = UserAccount(user_id=user1_id, email="adviser@school.edu", password_hash="hash", account_status="ACTIVE")
    other_user = UserAccount(user_id=user2_id, email="other@school.edu", password_hash="hash", account_status="ACTIVE")
    adviser_staff = AcademicStaff(staff_id="STAFF_ADV", user_id=user1_id, first_name="Maria", last_name="Adviser")
    other_staff = AcademicStaff(staff_id="STAFF_OTHER", user_id=user2_id, first_name="Pedro", last_name="Teacher")

    # Class Section
    class_8a = Class(class_id=1, section_name="Diamond", academic_year_id=1, academic_level_id=1, adviser_staff_id="STAFF_ADV")

    # Subjects & Subject Loads
    s1 = Subject(subject_id=101, subject_name="Mathematics 8", subject_codename="MATH8")
    s2 = Subject(subject_id=102, subject_name="Science 8", subject_codename="SCI8")
    load1 = SubjectLoad(subject_load_id=1, staff_id="STAFF_ADV", subject_id=101, class_id=1, academic_period_id=1, is_active_version=True, status="active")
    load2 = SubjectLoad(subject_load_id=2, staff_id="STAFF_OTHER", subject_id=102, class_id=1, academic_period_id=1, is_active_version=True, status="active")

    # Students (2 students)
    s1_id = uuid.uuid4()
    s2_id = uuid.uuid4()
    stud1 = Student(student_id=s1_id, student_lrn="123456789001", first_name="Ana", last_name="Reyes", gender="Female")
    stud2 = Student(student_id=s2_id, student_lrn="123456789002", first_name="Ben", last_name="Santos", gender="Male")

    sc1 = StudentClass(student_id=s1_id, class_id=1, academic_year_id=1, enrollment_status="enrolled")
    sc2 = StudentClass(student_id=s2_id, class_id=1, academic_year_id=1, enrollment_status="enrolled")

    # Period Grades for Student 1
    g1 = StudentPeriodGrade(student_id=s1_id, class_id=1, subject_id=101, academic_period_id=1, final_period_grade=92.0, is_finalized=True)
    g2 = StudentPeriodGrade(student_id=s1_id, class_id=1, subject_id=101, academic_period_id=2, final_period_grade=94.0, is_finalized=True)
    g3 = StudentPeriodGrade(student_id=s1_id, class_id=1, subject_id=102, academic_period_id=1, final_period_grade=88.0, is_finalized=True)
    g4 = StudentPeriodGrade(student_id=s1_id, class_id=1, subject_id=102, academic_period_id=2, final_period_grade=90.0, is_finalized=True)

    # Attendance Records
    att1 = AttendanceRecord(student_id=s1_id, class_id=1, date=date(2026, 8, 10), status="present")
    att2 = AttendanceRecord(student_id=s1_id, class_id=1, date=date(2026, 8, 11), status="present")
    att3 = AttendanceRecord(student_id=s1_id, class_id=1, date=date(2026, 8, 12), status="absent")

    db.add_all([
        ay, level, p1, p2,
        adviser_user, other_user, adviser_staff, other_staff,
        class_8a, s1, s2, load1, load2,
        stud1, stud2, sc1, sc2,
        g1, g2, g3, g4,
        att1, att2, att3,
    ])
    db.commit()

    return db, {
        "class_id": 1,
        "student1_id": s1_id,
        "student2_id": s2_id,
        "adviser_staff_id": "STAFF_ADV",
        "other_staff_id": "STAFF_OTHER",
    }


def test_export_student_sf9_structure():
    db, ctx = setup_test_db()
    stream = export_student_sf9(
        db=db,
        class_id=ctx["class_id"],
        student_id=ctx["student1_id"],
        staff_id=ctx["adviser_staff_id"],
    )
    assert stream is not None
    wb = openpyxl.load_workbook(stream)
    assert len(wb.sheetnames) == 1
    ws = wb.active
    assert ws is not None

    # Check headers and sections across full sheet
    cell_values = [ws.cell(row=r, column=c).value for r in range(1, ws.max_row + 1) for c in range(1, 8)]
    cell_text = " ".join([str(v) for v in cell_values if v is not None])

    assert "Department of Education" in cell_text
    assert "LEARNER'S PROGRESS REPORT CARD" in cell_text
    assert "Ana Reyes" in cell_text or "Reyes" in cell_text
    assert "123456789001" in cell_text
    assert "Grade 8 - Diamond" in cell_text
    assert "Mathematics 8" in cell_text
    assert "Science 8" in cell_text
    assert "REPORT ON LEARNER'S OBSERVED VALUES" in cell_text
    assert "Maka-Diyos" in cell_text
    assert "REPORT ON ATTENDANCE" in cell_text
    assert "CERTIFICATE OF TRANSFER" in cell_text


def test_export_student_sf9_unauthorized_staff():
    db, ctx = setup_test_db()
    with pytest.raises(Exception) as excinfo:
        export_student_sf9(
            db=db,
            class_id=ctx["class_id"],
            student_id=ctx["student1_id"],
            staff_id=ctx["other_staff_id"],
        )
    assert "not authorized" in str(excinfo.value) or "403" in str(excinfo.value)


def test_export_advisory_class_sf9_batch():
    db, ctx = setup_test_db()
    stream = export_advisory_class_sf9(
        db=db,
        class_id=ctx["class_id"],
        staff_id=ctx["adviser_staff_id"],
    )
    assert stream is not None
    wb = openpyxl.load_workbook(stream)
    # 2 enrolled students -> 2 worksheets in batch export
    assert len(wb.sheetnames) == 2


def test_api_export_sf9_endpoints():
    db, ctx = setup_test_db()

    app = FastAPI()
    app.include_router(classes_router, prefix="/api/v1/classes")

    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_staff_id] = lambda: ctx["adviser_staff_id"]

    client = TestClient(app)

    # 1. Single student export
    res1 = client.get(f"/api/v1/classes/teacher/advisory/{ctx['class_id']}/students/{ctx['student1_id']}/export-sf9")
    assert res1.status_code == 200
    assert "application/vnd.openxmlformats" in res1.headers["content-type"]
    wb1 = openpyxl.load_workbook(io.BytesIO(res1.content))
    assert len(wb1.sheetnames) == 1

    # 2. Batch advisory class export
    res2 = client.get(f"/api/v1/classes/teacher/advisory/{ctx['class_id']}/export-sf9-batch")
    assert res2.status_code == 200
    assert "application/vnd.openxmlformats" in res2.headers["content-type"]
    wb2 = openpyxl.load_workbook(io.BytesIO(res2.content))
    assert len(wb2.sheetnames) == 2

    # 3. SF9 data endpoint
    res3 = client.get(f"/api/v1/classes/teacher/advisory/{ctx['class_id']}/students/{ctx['student1_id']}/sf9-data")
    assert res3.status_code == 200
    data = res3.json()
    assert data["student"]["student_lrn"] == "123456789001"
    assert len(data["periods"]) == 2
    assert "attendance" in data
    assert len(data["attendance"]["months"]) == 11

    # 4. Word (.docx) single student export
    res4 = client.get(f"/api/v1/classes/teacher/advisory/{ctx['class_id']}/students/{ctx['student1_id']}/export-sf9-docx?term1_comment=Great+job")
    assert res4.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res4.headers["content-type"]
    assert len(res4.content) > 1000

    # 5. Word (.docx) batch export
    res5 = client.get(f"/api/v1/classes/teacher/advisory/{ctx['class_id']}/export-sf9-batch-docx")
    assert res5.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res5.headers["content-type"]
    assert len(res5.content) > 1000

