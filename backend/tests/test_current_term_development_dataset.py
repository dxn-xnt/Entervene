from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from app.ml.BuildCurrentTermDevelopmentDataset import (
    DATASET_FILE,
    FEATURE_COLUMNS,
    SUMMARY_FILE,
    TARGET_COLUMN,
    assert_no_target_leakage,
    assert_student_group_split,
    build_dataset,
    normalize_subject,
)


def _save_workbook(
    path: Path,
    *,
    subject: str = "MATHEMATICS",
    weights: tuple[int, int, int] = (20, 50, 30),
    student_name: str = "TEST, STUDENT",
    target: float = 92.0,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SUBJECT_Q1"
    ws.cell(5, 30, "SCHOOL YEAR")
    ws.cell(5, 33, "2025-2026")
    ws.cell(7, 6, "GRADE & SECTION: ")
    ws.cell(7, 11, "9-ARCHIMEDES")
    ws.cell(7, 29, "SUBJECT:")
    ws.cell(7, 33, subject)
    ws.cell(8, 2, "LEARNERS' NAMES")
    ws.cell(8, 6, f"WRITTEN WORKS ({weights[0]}%)")
    ws.cell(8, 13, f"PERFORMANCE TASKS ({weights[1]}%)")
    ws.cell(8, 20, f"QUARTERLY ASSESSMENT ({weights[2]}%)")
    ws.cell(8, 23, "Initial")
    ws.cell(8, 24, "Quarterly")

    labels = {
        6: 1, 7: 2, 8: 3, 9: 4, 10: "Total", 11: "PS", 12: "WS",
        13: 1, 14: 2, 15: 3, 16: 4, 17: "Total", 18: "PS", 19: "WS",
        20: 1, 21: "PS", 22: "WS", 23: "Grade", 24: "Grade",
    }
    for col, value in labels.items():
        ws.cell(9, col, value)
    ws.cell(10, 2, "HIGHEST POSSIBLE SCORE")
    for col in (6, 7, 8, 9, 13, 14, 15, 16):
        ws.cell(10, col, 10)
    ws.cell(10, 10, 40)
    ws.cell(10, 11, 100)
    ws.cell(10, 12, weights[0] / 100)
    ws.cell(10, 17, 40)
    ws.cell(10, 18, 100)
    ws.cell(10, 19, weights[1] / 100)
    ws.cell(10, 20, 40)
    ws.cell(10, 21, 100)
    ws.cell(10, 22, weights[2] / 100)

    ws.cell(11, 1, "MALE")
    ws.cell(12, 1, 1)
    ws.cell(12, 2, student_name)
    for offset, score in enumerate([10, 8, 6, 4]):
        ws.cell(12, 6 + offset, score)
    for offset, score in enumerate([9, 9, 8, 8]):
        ws.cell(12, 13 + offset, score)
    ws.cell(12, 20, 32)
    ws.cell(12, 24, target)
    wb.save(path)


def _read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def test_builds_same_period_development_dataset_and_summary(tmp_path: Path):
    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()
    _save_workbook(raw_dir / "MATH.xlsx", subject="MATHEMATICS", weights=(20, 50, 30), student_name="ONE, STUDENT")
    _save_workbook(raw_dir / "ICT.xlsx", subject="ICT", weights=(20, 60, 20), student_name="TWO, STUDENT")
    _save_workbook(raw_dir / "BAD_WEIGHT.xlsx", subject="SCIENCE", weights=(20, 40, 40), student_name="THREE, STUDENT")
    _save_workbook(raw_dir / "BAD_SUBJECT.xlsx", subject="FILIPINO", weights=(20, 50, 30), student_name="FOUR, STUDENT")

    summary = build_dataset(raw_dir, out_dir)
    rows = _read_rows(out_dir / DATASET_FILE)

    assert summary["target_column"] == TARGET_COLUMN
    assert summary["snapshot_row_count"] == 6
    assert summary["grading_weight_distribution"] == {"20/50/30": 3, "20/60/20": 3}
    assert summary["exclusion_counts"] == {"unsupported_subject": 3, "unsupported_weight_pattern": 3}
    assert summary["target_distribution"]["below_75_count"] == 0
    assert (out_dir / SUMMARY_FILE).exists()
    assert {row["target_source"] for row in rows} == {"final_quarterly_grade_column"}


def test_feature_schema_has_no_target_or_identity_leakage():
    forbidden = {TARGET_COLUMN, "raw_student_key", "student_period_key", "row_identity", "snapshot_fraction"}

    assert forbidden.isdisjoint(FEATURE_COLUMNS)
    assert_no_target_leakage([])


def test_student_group_split_keeps_all_snapshots_together(tmp_path: Path):
    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()
    for idx in range(8):
        _save_workbook(
            raw_dir / f"MATH_{idx}.xlsx",
            subject="MATHEMATICS",
            weights=(20, 50, 30),
            student_name=f"STUDENT{idx}, TEST",
        )

    build_dataset(raw_dir, out_dir)
    rows = _read_rows(out_dir / DATASET_FILE)
    assert_student_group_split(rows)
    per_student_splits: dict[str, set[str]] = {}
    for row in rows:
        per_student_splits.setdefault(row["raw_student_key"], set()).add(row["split"])
    assert all(len(splits) == 1 for splits in per_student_splits.values())


def test_subject_normalization_is_explicit_and_conservative():
    assert normalize_subject("Math 9") == "MATHEMATICS"
    assert normalize_subject("Science") == "SCIENCE"
    assert normalize_subject("Values Education") == "VALUES_EDUCATION"
    assert normalize_subject("Electronics") == "ELECTRONICS"
    assert normalize_subject("Filipino") is None
    assert normalize_subject("Araling Panlipunan") is None
    assert normalize_subject("Oral Communication") is None


def test_historical_qa_is_preserved_as_aggregate_not_fabricated(tmp_path: Path):
    raw_dir = tmp_path / "raw"
    out_dir = tmp_path / "out"
    raw_dir.mkdir()
    _save_workbook(raw_dir / "MATH.xlsx", subject="MATHEMATICS", weights=(20, 50, 30))

    build_dataset(raw_dir, out_dir)
    rows = _read_rows(out_dir / DATASET_FILE)

    assert {row["historical_exam_semantics"] for row in rows} == {"AGGREGATE_QA_AS_EXAM_ANALOGUE"}
    assert all("st1_percent" not in row for row in rows)
    assert all("st2_percent" not in row for row in rows)
    assert all("term_exam_percent" not in row for row in rows)
    early = {row["snapshot_fraction"]: row for row in rows}
    assert early["0.25"]["qa_available_activity_count"] == "0"
    assert early["0.75"]["qa_available_activity_count"] == "1"
