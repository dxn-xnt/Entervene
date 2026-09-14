from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from app.ml.BuildCurrentPeriodProjectionDataset import build_dataset
from app.ml.EClassRecordRawParser import EClassRecordRawParser


def _save_workbook(path: Path, *, weights=(40, 40, 20), malformed=False, over_hps=False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MATH_Q1"
    ws.cell(5, 30, "SCHOOL YEAR")
    ws.cell(5, 33, "2025-2026")
    ws.cell(7, 6, "GRADE & SECTION: ")
    ws.cell(7, 11, "9-ARCHIMEDES")
    ws.cell(7, 29, "SUBJECT:")
    ws.cell(7, 33, "MATHEMATICS")
    if malformed:
        ws.cell(8, 2, "LEARNERS' NAMES")
        ws.cell(12, 1, 1)
        ws.cell(12, 2, "TEST, STUDENT")
        wb.save(path)
        return

    ws.cell(8, 2, "LEARNERS' NAMES")
    ws.cell(8, 6, f"WRITTEN WORKS ({weights[0]}%)")
    ws.cell(8, 13, f"PERFORMANCE TASKS ({weights[1]}%)")
    ws.cell(8, 20, f"QUARTERLY ASSESSMENT ({weights[2]}%)")
    ws.cell(8, 23, "Initial")
    ws.cell(8, 24, "Quarterly")

    labels = {
        6: 1, 7: 2, 8: 3, 9: 4, 10: "Total", 11: "PS", 12: "WS",
        13: 1, 14: 2, 15: 3, 16: 4, 17: "Total", 18: "PS", 19: "WS",
        20: 1, 21: "PS", 22: "WS", 23: "Grade", 24: "Grade",
    }
    for col, value in labels.items():
        ws.cell(9, col, value)
    ws.cell(10, 2, "HIGHEST POSSIBLE SCORE")
    for col in (6, 7, 8, 9, 13, 14, 15, 16):
        ws.cell(10, col, 10)
    ws.cell(10, 10, 40)
    ws.cell(10, 11, 100)
    ws.cell(10, 12, weights[0] / 100)
    ws.cell(10, 17, 40)
    ws.cell(10, 18, 100)
    ws.cell(10, 19, weights[1] / 100)
    ws.cell(10, 20, 40)
    ws.cell(10, 21, 100)
    ws.cell(10, 22, weights[2] / 100)

    ws.cell(11, 1, "MALE")
    ws.cell(12, 1, 1)
    ws.cell(12, 2, "TEST, STUDENT")
    ww = [10, 8, 6, 4]
    pt = [9, 9, 8, 8]
    if over_hps:
        ww[0] = 12
    for offset, score in enumerate(ww):
        ws.cell(12, 6 + offset, score)
    ws.cell(12, 10, sum(ww))
    ws.cell(12, 11, sum(ww) / 40 * 100)
    ws.cell(12, 12, sum(ww) / 40 * weights[0])
    for offset, score in enumerate(pt):
        ws.cell(12, 13 + offset, score)
    ws.cell(12, 17, sum(pt))
    ws.cell(12, 18, sum(pt) / 40 * 100)
    ws.cell(12, 19, sum(pt) / 40 * weights[1])
    ws.cell(12, 20, 32)
    ws.cell(12, 21, 80)
    ws.cell(12, 22, weights[2] * 0.8)
    ws.cell(12, 23, 88)
    ws.cell(12, 24, 92)
    wb.save(path)


def test_one_student_period_produces_25_50_75_snapshots(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook)
    parser = EClassRecordRawParser()

    rows = parser.parse_workbook(workbook)

    assert [row["snapshot_fraction"] for row in rows] == [0.25, 0.5, 0.75]
    assert len({row["student_period_key"] for row in rows}) == 1
    assert rows[0]["target_final_period_grade"] == 92


def test_target_and_completed_summary_fields_do_not_appear_as_input_features(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook)
    row = EClassRecordRawParser().parse_workbook(workbook)[0]
    forbidden_input_names = {
        "initial_grade", "quarterly_grade", "final_grade", "transmuted_grade",
        "ww_ps", "pt_ps", "qa_ps", "ww_ws", "pt_ws", "qa_ws",
        "completed_quarter_ps", "completed_quarter_ws",
    }

    assert row["target_final_period_grade"] == 92
    assert forbidden_input_names.isdisjoint(row.keys())
    assert not any(key.endswith("_full_period_percent_audit") for key in row)


def test_future_activity_columns_do_not_enter_earlier_snapshots(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook)
    rows = EClassRecordRawParser().parse_workbook(workbook)
    by_fraction = {row["snapshot_fraction"]: row for row in rows}

    assert by_fraction[0.25]["ww_available_activity_count"] == 1
    assert by_fraction[0.25]["ww_points_earned_so_far"] == 10
    assert by_fraction[0.50]["ww_available_activity_count"] == 2
    assert by_fraction[0.50]["ww_points_earned_so_far"] == 18
    assert by_fraction[0.75]["ww_available_activity_count"] == 3
    assert by_fraction[0.75]["ww_points_earned_so_far"] == 24


def test_qa_is_not_exposed_in_early_snapshots(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook)
    rows = EClassRecordRawParser().parse_workbook(workbook)
    by_fraction = {row["snapshot_fraction"]: row for row in rows}

    assert by_fraction[0.25]["qa_available_activity_count"] == 0
    assert by_fraction[0.25]["qa_percent_so_far"] is None
    assert by_fraction[0.50]["qa_available_activity_count"] == 0
    assert by_fraction[0.75]["qa_available_activity_count"] == 1
    assert by_fraction[0.75]["qa_percent_so_far"] == 80


def test_malformed_sheet_reports_safely(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook, malformed=True)
    parser = EClassRecordRawParser()

    rows = parser.parse_workbook(workbook)

    assert rows == []
    assert parser.sheet_audits[0].status == "SKIPPED"
    assert "LAYOUT_NOT_FOUND" in parser.sheet_audits[0].reason


def test_grouping_keys_are_stable(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook)

    first = EClassRecordRawParser().parse_workbook(workbook)
    second = EClassRecordRawParser().parse_workbook(workbook)

    assert [row["raw_student_key"] for row in first] == [row["raw_student_key"] for row in second]
    assert [row["student_period_key"] for row in first] == [row["student_period_key"] for row in second]


@pytest.mark.parametrize("weights", [(40, 40, 20), (20, 60, 20), (30, 50, 20), (25, 45, 30)])
def test_component_weight_variations_are_preserved(tmp_path: Path, weights):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook, weights=weights)
    row = EClassRecordRawParser().parse_workbook(workbook)[0]

    assert row["ww_weight"] == weights[0]
    assert row["pt_weight"] == weights[1]
    assert row["qa_weight"] == weights[2]


def test_over_100_scores_are_reported_not_clipped(tmp_path: Path):
    workbook = tmp_path / "MATH 9-ARCHIMEDES.xlsx"
    _save_workbook(workbook, over_hps=True)
    parser = EClassRecordRawParser()
    rows = parser.parse_workbook(workbook)

    assert rows[0]["ww_percent_so_far"] == 120
    assert any(anomaly.anomaly_type == "ACTIVITY_SCORE_OVER_HPS" for anomaly in parser.anomalies)


def test_builder_writes_dataset_and_audit_files(tmp_path: Path):
    raw_dir = tmp_path / "raw"
    output_dir = tmp_path / "out"
    raw_dir.mkdir()
    _save_workbook(raw_dir / "MATH 9-ARCHIMEDES.xlsx")

    summary = build_dataset(raw_dir, output_dir)

    assert summary["snapshot_row_count"] == 3
    assert summary["usable_student_period_count"] == 1
    assert (output_dir / "current_period_projection_snapshots.csv").exists()
    assert (output_dir / "sheet_audit.csv").exists()
    with (output_dir / "current_period_projection_snapshots.csv").open(newline="", encoding="utf-8") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert len(csv_rows) == 3
