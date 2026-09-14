from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import uuid
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - exercised by CLI environment
    raise SystemExit("openpyxl is required to parse E-Class Record workbooks") from exc

COMPONENTS = ("ww", "pt", "qa")
SNAPSHOT_FRACTIONS = (0.25, 0.50, 0.75)
SNAPSHOT_POLICY_VERSION = "component-column-order-v1_qa-at-75-only"


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def _upper(value: Any) -> str:
    return _clean_text(value).upper()


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if math.isfinite(float(value)):
            return float(value)
        return None
    text = _clean_text(value).replace(",", "")
    if text == "":
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _stable_uuid(namespace_text: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, namespace_text))


def _normalize_subject(value: str) -> str:
    text = _upper(value)
    if not text:
        return "UNKNOWN"
    if "PHYSICS" in text:
        return "ADVANCED_PHYSICS"
    if "ELECTRONICS" in text:
        return "ELECTRONICS"
    if "CON CHEM" in text or "CHEM" in text or "SCIENCE" in text:
        return "SCIENCE"
    if "MATH" in text:
        return "MATHEMATICS"
    if "PRECAL" in text:
        return "PRE_CALCULUS"
    if "CREATIVE" in text and "TECH" in text:
        return "CREATIVE_TECHNOLOGY"
    if "VALUES" in text:
        return "VALUES_EDUCATION"
    if text == "ICT" or " ICT" in text:
        return "ICT"
    if (
        "MAPEH" in text
        or "PHYSICAL EDUCATION" in text
        or text in {"MUSIC", "ARTS", "P.E.", "PE", "HEALTH"}
    ):
        return "MAPEH"
    if "ORAL" in text and "COM" in text:
        return "ORAL_COMMUNICATION"
    if "PERDEV" in text or "PERSONAL DEVELOPMENT" in text:
        return "PERSONAL_DEVELOPMENT"
    if "ENGLISH" in text:
        return "ENGLISH"
    return re.sub(r"[^A-Z0-9]+", "_", text).strip("_") or "UNKNOWN"


def infer_file_metadata(path: Path) -> dict[str, Any]:
    name = path.name
    up = name.upper()
    grade_level = None
    grade_match = re.search(r"(?:GRADE[- ]?)?(7|8|9|10|11|12)\b", up)
    if grade_match:
        grade_level = int(grade_match.group(1))

    section = "UNKNOWN"
    for candidate in (
        "EINSTEIN", "SOCRATES", "ARCHIMEDES", "COPERNICUS", "NEWTON", "PLATO",
        "ARISTOTLE", "GALILEO", "CAMPOS", "ZARA", "DELMUNDO",
    ):
        if candidate in up:
            section = candidate
            break

    subject = "UNKNOWN"
    for marker, normalized in (
        ("PHYSICS", "ADVANCED_PHYSICS"),
        ("ELECTRONICS", "ELECTRONICS"),
        ("CON CHEM", "SCIENCE"),
        ("SCIENCE", "SCIENCE"),
        ("MATH", "MATHEMATICS"),
        ("PRECALCULUS", "PRE_CALCULUS"),
        ("CREATIVE TECH", "CREATIVE_TECHNOLOGY"),
        ("VALUES", "VALUES_EDUCATION"),
        ("ICT", "ICT"),
        ("MAPEH", "MAPEH"),
        ("ORAL COM", "ORAL_COMMUNICATION"),
        ("PERDEV", "PERSONAL_DEVELOPMENT"),
        ("ENGLISH", "ENGLISH"),
    ):
        if marker in up:
            subject = normalized
            break

    school_year = "UNKNOWN"
    if "2022-23" in up or "2022-2023" in up:
        school_year = "2022-2023"
    elif "2023 - 2024" in up or "2023-2024" in up:
        school_year = "2023-2024"
    elif "24-25" in up or "2024-2025" in up:
        school_year = "2024-2025"

    return {
        "grade_level": grade_level,
        "section": section,
        "subject": subject,
        "school_year": school_year,
    }


def period_sequence_from_sheet(sheet_name: str) -> int | None:
    up = sheet_name.upper()
    if "Q1" in up or "1ST" in up or "FIRST" in up:
        return 1
    if "Q2" in up or "2ND" in up or "SECOND" in up:
        return 2
    if "Q3" in up or "3RD" in up or "THIRD" in up:
        return 3
    if "Q4" in up or "4TH" in up or "FOURTH" in up:
        return 4
    return None


@dataclass(frozen=True)
class ComponentLayout:
    key: str
    label: str
    start_col: int
    activity_cols: tuple[int, ...]
    hps_by_col: dict[int, float]
    weight: float | None
    total_col: int | None = None
    ps_col: int | None = None
    ws_col: int | None = None


@dataclass(frozen=True)
class SheetLayout:
    header_row_idx: int
    label_row_idx: int
    hps_row_idx: int
    final_grade_col: int
    components: dict[str, ComponentLayout]


@dataclass
class RawAnomaly:
    source_file: str
    sheet_name: str
    row_number: int | None
    student_period_key: str | None
    component: str | None
    anomaly_type: str
    classification: str
    detail: str


@dataclass
class RemovedRow:
    source_file: str
    sheet_name: str
    row_number: int | None
    raw_student_name: str | None
    reason: str


@dataclass
class SheetAudit:
    source_file: str
    sheet_name: str
    status: str
    reason: str
    rows: int = 0
    snapshots: int = 0
    subject: str = "UNKNOWN"
    grade_level: int | None = None
    section: str = "UNKNOWN"
    school_year: str = "UNKNOWN"
    period_sequence: int | None = None
    weights: dict[str, float | None] = field(default_factory=dict)


class EClassRecordRawParser:
    """Parse DepEd E-Class Record workbooks into leakage-safe partial snapshots.

    The parser reads raw individual activity columns and HPS values. Completed-quarter
    PS, WS, Initial Grade, and Quarterly Grade columns are used only for audit/target
    extraction and never as input features.
    """

    def __init__(self, snapshot_fractions: Iterable[float] = SNAPSHOT_FRACTIONS):
        self.snapshot_fractions = tuple(snapshot_fractions)
        self.anomalies: list[RawAnomaly] = []
        self.removed_rows: list[RemovedRow] = []
        self.sheet_audits: list[SheetAudit] = []

    def parse_directory(self, raw_dir: Path) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for workbook_path in sorted(raw_dir.glob("*.xlsx")):
            rows.extend(self.parse_workbook(workbook_path))
        return rows

    def parse_workbook(self, workbook_path: Path) -> list[dict[str, Any]]:
        try:
            wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
        except Exception as exc:
            self.sheet_audits.append(SheetAudit(
                source_file=workbook_path.name,
                sheet_name="<workbook>",
                status="SKIPPED",
                reason=f"WORKBOOK_OPEN_FAILED: {exc}",
            ))
            return []

        snapshots: list[dict[str, Any]] = []
        for sheet_name in wb.sheetnames:
            period_sequence = period_sequence_from_sheet(sheet_name)
            if period_sequence is None:
                self.sheet_audits.append(SheetAudit(
                    source_file=workbook_path.name,
                    sheet_name=sheet_name,
                    status="SKIPPED",
                    reason="NOT_A_GRADING_PERIOD_SHEET",
                ))
                continue
            snapshots.extend(self.parse_sheet(workbook_path, wb[sheet_name], sheet_name, period_sequence))
        return snapshots

    def parse_sheet(self, workbook_path: Path, worksheet: Any, sheet_name: str, period_sequence: int) -> list[dict[str, Any]]:
        all_rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        metadata = self._metadata_from_sheet(workbook_path, all_rows)
        metadata["period_sequence"] = period_sequence
        metadata["historical_period_label"] = sheet_name

        try:
            layout = self._find_layout(all_rows)
        except ValueError as exc:
            self.sheet_audits.append(SheetAudit(
                source_file=workbook_path.name,
                sheet_name=sheet_name,
                status="SKIPPED",
                reason=str(exc),
                subject=metadata["subject"],
                grade_level=metadata["grade_level"],
                section=metadata["section"],
                school_year=metadata["school_year"],
                period_sequence=period_sequence,
            ))
            return []

        student_rows = self._student_rows(all_rows, layout.hps_row_idx + 1)
        snapshots: list[dict[str, Any]] = []
        rows_with_target = 0
        rows_without_target = 0
        for row_idx, name, row in student_rows:
            target = _to_float(row[layout.final_grade_col] if layout.final_grade_col < len(row) else None)
            if target is None:
                rows_without_target += 1
                self.removed_rows.append(RemovedRow(workbook_path.name, sheet_name, row_idx + 1, name, "MISSING_FINAL_TARGET_GRADE"))
                continue
            rows_with_target += 1
            student_rows_for_period = self._snapshots_for_student(
                workbook_path=workbook_path,
                sheet_name=sheet_name,
                row_number=row_idx + 1,
                raw_student_name=name,
                row=row,
                layout=layout,
                metadata=metadata,
                target_final_period_grade=target,
            )
            if not any(snapshot.get("has_any_input_evidence") for snapshot in student_rows_for_period):
                self.removed_rows.append(RemovedRow(
                    workbook_path.name,
                    sheet_name,
                    row_idx + 1,
                    name,
                    "NO_RAW_ACTIVITY_EVIDENCE",
                ))
                continue
            snapshots.extend(student_rows_for_period)

        component_count = sum(1 for component in COMPONENTS if layout.components.get(component))
        missing_weights = [key for key, component in layout.components.items() if component.weight is None]
        if not snapshots:
            status = "SKIPPED"
            reason = "NO_STUDENT_ROWS_WITH_FINAL_TARGET" if rows_with_target == 0 else "NO_SNAPSHOTS_CREATED"
        elif component_count < 3 or missing_weights or rows_without_target:
            status = "PARTIAL"
            reasons = []
            if component_count < 3:
                reasons.append("MISSING_COMPONENT_LAYOUT")
            if missing_weights:
                reasons.append(f"MISSING_COMPONENT_WEIGHT:{','.join(missing_weights)}")
            if rows_without_target:
                reasons.append(f"REMOVED_ROWS_WITHOUT_TARGET:{rows_without_target}")
            reason = ";".join(reasons) or "PARTIAL_PARSE"
        else:
            status = "PARSED"
            reason = "OK"

        self.sheet_audits.append(SheetAudit(
            source_file=workbook_path.name,
            sheet_name=sheet_name,
            status=status,
            reason=reason,
            rows=rows_with_target,
            snapshots=len(snapshots),
            subject=metadata["subject"],
            grade_level=metadata["grade_level"],
            section=metadata["section"],
            school_year=metadata["school_year"],
            period_sequence=period_sequence,
            weights={key: layout.components[key].weight if key in layout.components else None for key in COMPONENTS},
        ))
        return snapshots

    def _metadata_from_sheet(self, workbook_path: Path, rows: list[tuple[Any, ...]]) -> dict[str, Any]:
        metadata = infer_file_metadata(workbook_path)
        subject_from_sheet = None
        grade_section = None
        school_year = None
        for row in rows[:15]:
            for idx, cell in enumerate(row):
                cell_text = _upper(cell)
                if cell_text == "SUBJECT:":
                    subject_from_sheet = self._next_text(row, idx)
                elif cell_text == "GRADE & SECTION:":
                    grade_section = self._next_text(row, idx)
                elif cell_text == "SCHOOL YEAR":
                    school_year = self._next_text(row, idx)
        if subject_from_sheet:
            metadata["subject"] = _normalize_subject(subject_from_sheet)
        if grade_section:
            grade_match = re.search(r"(?:GRADE\s*)?(7|8|9|10|11|12|VII|VIII|IX|X|XI|XII)\b", grade_section.upper())
            roman = {"VII": 7, "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}
            if grade_match:
                token = grade_match.group(1)
                metadata["grade_level"] = int(token) if token.isdigit() else roman.get(token, metadata.get("grade_level"))
            for part in re.split(r"[-\s]+", grade_section.upper()):
                cleaned = re.sub(r"[^A-Z]", "", part)
                if cleaned and cleaned not in {"GRADE", "VII", "VIII", "IX", "X", "XI", "XII"}:
                    if cleaned not in {"7", "8", "9", "10", "11", "12"}:
                        metadata["section"] = cleaned
        if school_year:
            metadata["school_year"] = school_year
        if metadata.get("grade_level") is None:
            metadata["grade_level"] = "UNKNOWN"
        return metadata

    @staticmethod
    def _next_text(row: tuple[Any, ...], idx: int) -> str | None:
        for next_cell in row[idx + 1: idx + 8]:
            text = _clean_text(next_cell)
            if text:
                return text
        return None

    def _find_layout(self, rows: list[tuple[Any, ...]]) -> SheetLayout:
        header_idx = None
        for idx, row in enumerate(rows[:20]):
            upper_cells = [_upper(cell) for cell in row]
            if any("WRITTEN" in cell for cell in upper_cells) and any("PERFORMANCE" in cell for cell in upper_cells) and any("QUARTERLY ASSESSMENT" in cell for cell in upper_cells):
                header_idx = idx
                break
        if header_idx is None:
            raise ValueError("LAYOUT_NOT_FOUND: component header row missing")

        label_idx = None
        hps_idx = None
        for idx in range(header_idx + 1, min(header_idx + 5, len(rows))):
            upper_cells = [_upper(cell) for cell in rows[idx]]
            if sum(1 for cell in upper_cells if cell == "PS") >= 2 and any(cell == "TOTAL" for cell in upper_cells):
                label_idx = idx
            if any("HIGHEST POSSIBLE SCORE" in cell for cell in upper_cells):
                hps_idx = idx
        if label_idx is None:
            raise ValueError("LAYOUT_NOT_FOUND: activity label row missing")
        if hps_idx is None:
            raise ValueError("LAYOUT_NOT_FOUND: HPS row missing")

        header_row = rows[header_idx]
        component_starts: list[tuple[str, int, str]] = []
        for col, cell in enumerate(header_row):
            text = _upper(cell)
            if "WRITTEN" in text:
                component_starts.append(("ww", col, _clean_text(cell)))
            elif "PERFORMANCE" in text:
                component_starts.append(("pt", col, _clean_text(cell)))
            elif "QUARTERLY ASSESSMENT" in text:
                component_starts.append(("qa", col, _clean_text(cell)))
        component_starts.sort(key=lambda item: item[1])
        if len(component_starts) < 2:
            raise ValueError("LAYOUT_NOT_FOUND: fewer than two components found")

        label_row = rows[label_idx]
        hps_row = rows[hps_idx]
        grade_cols = [col for col, cell in enumerate(label_row) if "GRADE" in _upper(cell)]
        if not grade_cols:
            raise ValueError("LAYOUT_NOT_FOUND: final grade column missing")
        final_grade_col = grade_cols[1] if len(grade_cols) >= 2 else grade_cols[-1]

        components: dict[str, ComponentLayout] = {}
        for pos, (key, start_col, label) in enumerate(component_starts):
            next_start = component_starts[pos + 1][1] if pos + 1 < len(component_starts) else final_grade_col
            section_cols = range(start_col, next_start)
            total_col = ps_col = ws_col = None
            activity_cols: list[int] = []
            hps_by_col: dict[int, float] = {}
            for col in section_cols:
                label_text = _upper(label_row[col] if col < len(label_row) else None)
                hps_value = _to_float(hps_row[col] if col < len(hps_row) else None)
                if label_text == "TOTAL":
                    total_col = col
                    continue
                if label_text == "PS":
                    ps_col = col
                    continue
                if label_text == "WS":
                    ws_col = col
                    continue
                if hps_value is not None and hps_value > 0 and self._is_activity_label(label_text):
                    activity_cols.append(col)
                    hps_by_col[col] = hps_value
            weight = self._weight_from_label(label)
            if weight is None and ws_col is not None and ws_col < len(hps_row):
                ws_hps = _to_float(hps_row[ws_col])
                if ws_hps is not None and 0 < ws_hps <= 1:
                    weight = round(ws_hps * 100, 4)
                elif ws_hps is not None and 1 < ws_hps <= 100:
                    weight = round(ws_hps, 4)
            components[key] = ComponentLayout(
                key=key,
                label=label,
                start_col=start_col,
                activity_cols=tuple(activity_cols),
                hps_by_col=hps_by_col,
                weight=weight,
                total_col=total_col,
                ps_col=ps_col,
                ws_col=ws_col,
            )
        return SheetLayout(header_idx, label_idx, hps_idx, final_grade_col, components)

    @staticmethod
    def _is_activity_label(label_text: str) -> bool:
        if label_text == "":
            return True
        return bool(re.fullmatch(r"\d+(?:\.0+)?", label_text))

    @staticmethod
    def _weight_from_label(label: str) -> float | None:
        match = re.search(r"\((\d+(?:\.\d+)?)\s*%\)", label)
        if match:
            return float(match.group(1))
        return None

    def _student_rows(self, rows: list[tuple[Any, ...]], start_idx: int) -> list[tuple[int, str, tuple[Any, ...]]]:
        result: list[tuple[int, str, tuple[Any, ...]]] = []
        in_students = False
        for row_idx in range(start_idx, min(len(rows), start_idx + 120)):
            row = rows[row_idx]
            first = _upper(row[0] if len(row) > 0 else None)
            second = _clean_text(row[1] if len(row) > 1 else None)
            if first in {"MALE", "FEMALE"} or _upper(second) in {"MALE", "FEMALE"}:
                in_students = True
                continue
            if not in_students and not _to_float(row[0] if len(row) > 0 else None):
                continue
            if self._looks_like_student_row(row):
                name = self._student_name_from_row(row)
                if name:
                    result.append((row_idx, name, row))
        return result

    @staticmethod
    def _looks_like_student_row(row: tuple[Any, ...]) -> bool:
        first = _to_float(row[0] if len(row) > 0 else None)
        name = EClassRecordRawParser._student_name_from_row(row)
        return first is not None and name is not None

    @staticmethod
    def _student_name_from_row(row: tuple[Any, ...]) -> str | None:
        for idx in (1, 2, 3, 4):
            if idx >= len(row):
                continue
            text = _clean_text(row[idx])
            up = text.upper()
            if text and not up in {"MALE", "FEMALE"} and not re.fullmatch(r"\d+(?:\.0+)?", text):
                return text
        return None

    def _snapshots_for_student(
        self,
        *,
        workbook_path: Path,
        sheet_name: str,
        row_number: int,
        raw_student_name: str,
        row: tuple[Any, ...],
        layout: SheetLayout,
        metadata: dict[str, Any],
        target_final_period_grade: float,
    ) -> list[dict[str, Any]]:
        source_file = workbook_path.name
        learner_identity = self._learner_identity(source_file, metadata, raw_student_name, row_number)
        raw_student_key = _stable_uuid(f"eclass-learner::{learner_identity}")
        student_period_identity = (
            f"{source_file}::{metadata.get('section')}::{learner_identity}::"
            f"{metadata['school_year']}::{metadata['subject']}::{metadata['period_sequence']}::{sheet_name}"
        )
        student_period_key = _stable_uuid(f"eclass-student-period::{student_period_identity}")

        full_component_percents: dict[str, float | None] = {}
        for key, component in layout.components.items():
            earned, possible, used_count, total_count, over_hps_count = self._component_points(row, component, len(component.activity_cols))
            percent = self._percent(earned, possible)
            full_component_percents[key] = percent
            if over_hps_count:
                self.anomalies.append(RawAnomaly(
                    source_file, sheet_name, row_number, student_period_key, key,
                    "ACTIVITY_SCORE_OVER_HPS", "LEGITIMATE_BONUS_OR_EXTRA_CREDIT",
                    f"{over_hps_count} activity score(s) exceed HPS; earned={earned}, possible={possible}",
                ))
            if percent is not None and percent > 100:
                classification = "LEGITIMATE_BONUS_OR_EXTRA_CREDIT" if over_hps_count else "UNRESOLVED"
                self.anomalies.append(RawAnomaly(
                    source_file, sheet_name, row_number, student_period_key, key,
                    "RECALCULATED_COMPONENT_PERCENT_OVER_100", classification,
                    f"full raw component percent={percent:.4f}; no clipping applied",
                ))
            reported_ps = _to_float(row[component.ps_col] if component.ps_col is not None and component.ps_col < len(row) else None)
            if reported_ps is not None and reported_ps > 100:
                classification = "LEGITIMATE_BONUS_OR_EXTRA_CREDIT" if over_hps_count or (percent is not None and percent > 100) else "UNRESOLVED"
                self.anomalies.append(RawAnomaly(
                    source_file, sheet_name, row_number, student_period_key, key,
                    "REPORTED_PS_OVER_100", classification,
                    f"reported completed-quarter PS={reported_ps}; not used as input",
                ))

        snapshots = []
        for fraction in self.snapshot_fractions:
            record: dict[str, Any] = {
                "student_period_key": student_period_key,
                "raw_student_key": raw_student_key,
                "source_file": source_file,
                "sheet_name": sheet_name,
                "row_number": row_number,
                "raw_student_name": raw_student_name,
                "school_year": metadata["school_year"],
                "grade_level": metadata["grade_level"],
                "section": metadata["section"],
                "subject": metadata["subject"],
                "historical_period_label": metadata["historical_period_label"],
                "period_sequence_audit": metadata["period_sequence"],
                "snapshot_fraction": fraction,
                "snapshot_policy_version": SNAPSHOT_POLICY_VERSION,
                "target_final_period_grade": target_final_period_grade,
                "target_source": "final_quarterly_grade_column",
                "attendance_available": False,
            }
            observed_weight = 0.0
            weighted_score_sum = 0.0
            available_activity_count = 0
            total_activity_count = 0
            for key in COMPONENTS:
                component = layout.components.get(key)
                if component is None:
                    self._empty_component_fields(record, key)
                    continue
                reveal_count = self._revealed_activity_count(key, len(component.activity_cols), fraction)
                earned, possible, used_count, total_count, over_hps_count = self._component_points(row, component, reveal_count)
                percent = self._percent(earned, possible)
                weighted = percent * (component.weight or 0.0) / 100.0 if percent is not None and component.weight is not None else None
                record[f"{key}_weight"] = component.weight
                record[f"{key}_available_activity_count"] = used_count
                record[f"{key}_total_activity_count"] = total_count
                record[f"{key}_coverage_ratio"] = used_count / total_count if total_count else None
                record[f"{key}_points_earned_so_far"] = earned if used_count else None
                record[f"{key}_points_possible_so_far"] = possible if used_count else None
                record[f"{key}_percent_so_far"] = percent
                record[f"{key}_weighted_score_so_far"] = weighted
                record[f"{key}_has_evidence"] = used_count > 0
                record[f"{key}_score_over_hps_count_so_far"] = over_hps_count
                if used_count and percent is not None and component.weight is not None:
                    observed_weight += component.weight
                    weighted_score_sum += weighted or 0.0
                available_activity_count += used_count
                total_activity_count += total_count
            record["observed_component_weight_sum"] = observed_weight
            record["overall_weighted_score_so_far"] = weighted_score_sum if observed_weight else None
            record["overall_partial_percent"] = (weighted_score_sum / observed_weight * 100.0) if observed_weight else None
            record["overall_available_activity_count"] = available_activity_count
            record["overall_total_activity_count"] = total_activity_count
            record["overall_activity_coverage_ratio"] = available_activity_count / total_activity_count if total_activity_count else None
            record["has_any_input_evidence"] = available_activity_count > 0 and observed_weight > 0
            record["snapshot_trace_json"] = json.dumps(self._snapshot_trace(layout, fraction), sort_keys=True)
            snapshots.append(record)
        return snapshots

    @staticmethod
    def _learner_identity(source_file: str, metadata: dict[str, Any], raw_student_name: str, row_number: int) -> str:
        normalized_name = re.sub(r"[^a-z0-9]", "", raw_student_name.lower())
        if normalized_name and not re.fullmatch(r"(?:male|female)\d+", normalized_name):
            return f"name::{normalized_name}"
        return "local-row::{}::{}::{}::{}::{}".format(
            source_file,
            metadata.get("school_year"),
            metadata.get("section"),
            metadata.get("subject"),
            row_number,
        )

    @staticmethod
    def _revealed_activity_count(component_key: str, total_count: int, fraction: float) -> int:
        if total_count <= 0:
            return 0
        if component_key == "qa":
            return total_count if fraction >= 0.75 else 0
        return max(1, min(total_count, math.ceil(total_count * fraction)))

    @staticmethod
    def _component_points(row: tuple[Any, ...], component: ComponentLayout, reveal_count: int) -> tuple[float, float, int, int, int]:
        earned = 0.0
        possible = 0.0
        used = 0
        over_hps = 0
        total_count = len(component.activity_cols)
        for col in component.activity_cols[:reveal_count]:
            hps = component.hps_by_col[col]
            score = _to_float(row[col] if col < len(row) else None)
            if score is None:
                continue
            earned += score
            possible += hps
            used += 1
            if hps > 0 and score > hps:
                over_hps += 1
        return earned, possible, used, total_count, over_hps

    @staticmethod
    def _percent(earned: float, possible: float) -> float | None:
        if possible <= 0:
            return None
        return earned / possible * 100.0

    @staticmethod
    def _empty_component_fields(record: dict[str, Any], key: str) -> None:
        for suffix in (
            "weight", "available_activity_count", "total_activity_count", "coverage_ratio",
            "points_earned_so_far", "points_possible_so_far", "percent_so_far",
            "weighted_score_so_far", "has_evidence", "score_over_hps_count_so_far",
        ):
            record[f"{key}_{suffix}"] = None

    @staticmethod
    def _snapshot_trace(layout: SheetLayout, fraction: float) -> dict[str, Any]:
        trace = {"policy": SNAPSHOT_POLICY_VERSION, "fraction": fraction, "components": {}}
        for key, component in layout.components.items():
            count = EClassRecordRawParser._revealed_activity_count(key, len(component.activity_cols), fraction)
            trace["components"][key] = {
                "activity_columns_1_based": [col + 1 for col in component.activity_cols[:count]],
                "total_activity_columns_1_based": [col + 1 for col in component.activity_cols],
                "excluded_completed_summary_columns_1_based": [col + 1 for col in (component.total_col, component.ps_col, component.ws_col) if col is not None],
            }
        return trace


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = list(rows[0].keys())
    extra = sorted({key for row in rows for key in row.keys()} - set(fieldnames))
    fieldnames.extend(extra)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def records_sha256(rows: list[dict[str, Any]]) -> str:
    payload = json.dumps(rows, sort_keys=True, default=str, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def dataclass_rows(items: Iterable[Any]) -> list[dict[str, Any]]:
    rows = []
    for item in items:
        row = asdict(item)
        for key, value in list(row.items()):
            if isinstance(value, dict):
                row[key] = json.dumps(value, sort_keys=True)
        rows.append(row)
    return rows



