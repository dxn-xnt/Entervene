"""Offline, whole-workbook evaluation of the 4H official-target JHS candidate."""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import joblib
import numpy as np
import openpyxl
import pandas as pd
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

from app.ml.BuildCurrentTermDevelopmentDataset import FEATURE_COLUMNS
from app.ml.BuildCurrentTermOfficialTargetCandidate import (
    CANDIDATE_FILE,
    OUTPUT_DIR as DATA_DIR,
    RAW_DIR,
    build_candidate,
)
from app.ml.TrainCurrentTermDevelopmentModel import (
    RANDOM_STATE,
    TARGET_COLUMN,
    linear_pipeline,
    rf_pipeline,
)

BACKEND = Path(__file__).resolve().parents[2]
MODEL_NAME = "entervene_current_term_official_target_rf_candidate"
OUTPUT_DIR = BACKEND / "data" / "experiments" / "official_target_rf_4i"
MODEL_PATH = BACKEND / "data" / "models" / f"{MODEL_NAME}.joblib"
EXPECTED_DATA_HASH = "789879391162415bf2df25630e7bd65be0838a0b73b3e3cc62579e9e0019f5a1"
EXPECTED_LINEAGE_HASH = "9fc13be41410adffcf6ea9048f7e818d1674d3796bd1c75654bf42e95161fb86"
RUNTIME_FILES = {
    "CLASSRECORD. 7 - ARISTOTLE.xlsx", "CLASSRECORD. 7 -GALILEO.xlsx",
    "MATH 10 - EINSTEIN (2023 - 2024).xlsx",
    "MATH 10 - SOCRATES (2023 - 2024).xlsx",
    "MATH 10 -EINSTEIN 24-25.xlsx",
}
METHODS = ("rf", "global_mean", "subject_mean", "grade_subject_mean", "linear", "academic")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def measure(actual, predicted) -> dict:
    y = np.asarray(actual, dtype=float)
    p = np.asarray(predicted, dtype=float)
    if not len(y) or not np.isfinite(p).all():
        raise ValueError("Empty or non-finite evaluation")
    error = p - y
    return {
        "n": len(y), "mae": round(float(mean_absolute_error(y, p)), 4),
        "rmse": round(float(math.sqrt(mean_squared_error(y, p))), 4),
        "r2": round(float(r2_score(y, p)), 4) if len(y) > 1 else None,
        "mean_signed_error": round(float(error.mean()), 4),
        "within_1_pct": round(float(np.mean(np.abs(error) <= 1) * 100), 2),
        "within_2_pct": round(float(np.mean(np.abs(error) <= 2) * 100), 2),
        "within_3_pct": round(float(np.mean(np.abs(error) <= 3) * 100), 2),
    }


def verify_data() -> tuple[pd.DataFrame, dict]:
    rebuilt = build_candidate()
    data_path = DATA_DIR / CANDIDATE_FILE
    lineage_path = DATA_DIR / "target_lineage.csv"
    if rebuilt["candidate_sha256"] != EXPECTED_DATA_HASH or rebuilt["lineage_sha256"] != EXPECTED_LINEAGE_HASH:
        raise ValueError("Task 4H candidate hashes changed; stop before training")
    data = pd.read_csv(data_path)
    lineage = pd.read_csv(lineage_path)
    if (len(data), data.student_period_key.nunique(), data.raw_student_key.nunique(), data.source_file.nunique()) != (5595, 1865, 467, 14):
        raise ValueError("Task 4H counts changed; stop before training")
    if len(lineage) != 1865 or not lineage.resolution_status.eq("RESOLVED").all():
        raise ValueError("Official target resolution changed")
    if not data.target_source.eq("OFFICIAL_WORKBOOK_GRADE").all():
        raise ValueError("Mixed target source detected")
    if sorted(data.snapshot_fraction.unique()) != [0.25, 0.5, 0.75]:
        raise ValueError("Unexpected snapshot stages")
    if not data.groupby("student_period_key").size().eq(3).all():
        raise ValueError("Each period must have three snapshots")
    if (data.groupby("raw_student_key").source_file.nunique() > 1).any():
        raise ValueError("Student crosses workbooks")
    if (data.groupby("student_period_key")[TARGET_COLUMN].nunique() > 1).any():
        raise ValueError("Inconsistent target within period")
    schema = json.loads((BACKEND / "data/models/entervene_current_term_development_rf_v3_feature_schema.json").read_text())
    features = list(FEATURE_COLUMNS)
    if len(features) != 31 or features != schema["feature_columns"]:
        raise ValueError("Original V3 feature schema changed")
    forbidden = {"raw_student_key", "student_period_key", "source_file", "sheet_name", "row_number", "initial_grade", "quarterly_grade", TARGET_COLUMN}
    if forbidden.intersection(features):
        raise ValueError("PII, source, or outcome entered model features")
    if not data.loc[data.snapshot_fraction.lt(0.75), "qa_has_evidence"].eq(0).all():
        raise ValueError("Future QA evidence in early snapshot")
    return data, {"data_sha256": sha256(data_path), "lineage_sha256": sha256(lineage_path),
                  "feature_schema_sha256": hashlib.sha256(json.dumps(features, separators=(",", ":")).encode()).hexdigest(),
                  "periods": 1865, "students": 467, "snapshots": 5595, "workbooks": 14}


def folds(data: pd.DataFrame) -> list[dict]:
    result = []
    for filename in sorted(data.source_file.unique()):
        test = data[data.source_file.eq(filename)]
        train = data[~data.source_file.eq(filename)]
        if set(test.raw_student_key) & set(train.raw_student_key):
            raise ValueError("Student leakage across workbook fold")
        subject = test.subject.unique()
        if len(subject) != 1:
            raise ValueError("Workbook has multiple canonical subjects")
        grade = int(test.grade_level.iloc[0])
        subject_train = train[train.subject.eq(subject[0])]
        grade_subject_train = subject_train[subject_train.grade_level.eq(grade)]
        result.append({"file": filename, "subject": subject[0], "grade": grade,
                       "periods": test.student_period_key.nunique(), "students": test.raw_student_key.nunique(),
                       "class": "SEEN_SUBJECT_HOLDOUT" if len(subject_train) else "UNSEEN_SUBJECT_HOLDOUT",
                       "subject_train_periods": subject_train.student_period_key.nunique(),
                       "grade_subject_train_periods": grade_subject_train.student_period_key.nunique(),
                       "runtime_exact_match": filename in RUNTIME_FILES})
    return result


def historical_tables(files: list[str]) -> dict[str, list[tuple[float, float]]]:
    tables = {}
    for filename in files:
        workbook = openpyxl.load_workbook(RAW_DIR / filename, read_only=True, data_only=True)
        try:
            sheet = workbook["DO NOT DELETE"]
            rows = [(float(sheet.cell(row, 7).value), float(sheet.cell(row, 10).value)) for row in range(2, 43)]
            if len(rows) != 41 or rows != sorted(rows) or rows[0][0] != 0 or rows[-1][0] != 100:
                raise ValueError(f"Invalid historical lookup: {filename}")
            tables[filename] = rows
        finally:
            workbook.close()
    return tables


def historical_lookup_parity(data: pd.DataFrame, tables: dict[str, list[tuple[float, float]]]) -> dict:
    """Check the source lookup against cached AI and AJ, without using AI as a feature."""
    periods = data[data.snapshot_fraction.eq(0.75)]
    checked = mismatched = missing = 0
    for filename, source in periods.groupby("source_file"):
        workbook = openpyxl.load_workbook(RAW_DIR / filename, read_only=True, data_only=True)
        try:
            table = tables[filename]
            for sheet_name, sheet_rows in source.groupby("sheet_name"):
                wanted = set(int(number) for number in sheet_rows.row_number)
                start, end = min(wanted), max(wanted)
                for number, values in enumerate(workbook[sheet_name].iter_rows(
                        min_row=start, max_row=end, min_col=35, max_col=36, values_only=True), start):
                    if number not in wanted:
                        continue
                    initial, official = values
                    if not isinstance(initial, (int, float)) or not isinstance(official, (int, float)):
                        missing += 1
                        continue
                    checked += 1
                    expected = next(grade for minimum, grade in reversed(table) if round(float(initial), 2) >= minimum)
                    mismatched += expected != float(official)
        finally:
            workbook.close()
    return {"cached_initial_and_official_checked": checked, "missing": missing,
            "historical_lookup_mismatches": mismatched}


def academic_prediction(row, table: list[tuple[float, float]]) -> float:
    components = ("ww", "pt", "qa")
    if any(float(row[f"{part}_has_evidence"]) != 1 for part in components):
        return float("nan")
    initial = round(sum(float(row[f"{part}_weight"]) * float(row[f"{part}_percent_so_far"]) / 100
                        for part in components), 2)
    return next(grade for minimum, grade in reversed(table) if initial >= minimum)


def evaluate(data: pd.DataFrame, plan: list[dict], tables: dict[str, list[tuple[float, float]]]) -> pd.DataFrame:
    features = list(FEATURE_COLUMNS)
    rows = []
    for number, fold in enumerate(plan, 1):
        filename = fold["file"]
        train = data[~data.source_file.eq(filename)]
        test = data[data.source_file.eq(filename)].copy()
        rf = rf_pipeline(features)
        linear = linear_pipeline(features)
        rf.fit(train[features], train[TARGET_COLUMN])
        linear.fit(train[features], train[TARGET_COLUMN])
        # sklearn's parallel prediction sums tree outputs in completion order;
        # serial prediction makes the saved fold outputs byte-reproducible.
        rf.named_steps["regressor"].n_jobs = 1
        subject_means = train.groupby("subject")[TARGET_COLUMN].mean().to_dict()
        grade_subject_means = train.groupby(["grade_level", "subject"])[TARGET_COLUMN].mean().to_dict()
        predictions = {
            "rf": rf.predict(test[features]),
            "linear": linear.predict(test[features]),
            "global_mean": np.repeat(train[TARGET_COLUMN].mean(), len(test)),
            "subject_mean": test.subject.map(subject_means).to_numpy(),
            "grade_subject_mean": np.array([grade_subject_means.get((row.grade_level, row.subject), np.nan)
                                             for row in test.itertuples()], dtype=float),
            "academic": np.array([academic_prediction(row, tables[filename]) for _, row in test.iterrows()]),
        }
        for i, row in enumerate(test.itertuples()):
            output = {"fold": filename, "fold_class": fold["class"],
                      "runtime_exact_match": fold["runtime_exact_match"],
                      "pseudonymous_period_id": hashlib.sha256(row.student_period_key.encode()).hexdigest(),
                      "snapshot_fraction": row.snapshot_fraction, "subject": row.subject,
                      "grade_level": int(row.grade_level), "school_year": row.school_year,
                      "weight_pattern": row.weight_pattern, "actual": row.target_final_period_grade}
            output.update({method: float(values[i]) for method, values in predictions.items()})
            rows.append(output)
        print(f"fold {number}/{len(plan)}: {filename}", flush=True)
    return pd.DataFrame(rows)


def summarize(predictions: pd.DataFrame, plan: list[dict]) -> dict:
    seen = predictions[predictions.fold_class.eq("SEEN_SUBJECT_HOLDOUT")]
    runtime = seen[seen.runtime_exact_match]
    unseen = predictions[predictions.fold_class.eq("UNSEEN_SUBJECT_HOLDOUT")]
    out: dict = {"folds": plan, "views": {}, "subgroups": {}, "risk": {}}
    for label, frame in (("seen_historical", seen), ("runtime_exact", runtime), ("unseen_subject", unseen)):
        out["views"][label] = {"periods": len(frame) // 3, "workbooks": frame.fold.nunique(),
                               "students_note": "Students unique by source workbook"}
        for stage_label, stage in (("all_snapshots", frame), ("25_pct", frame[frame.snapshot_fraction.eq(0.25)]),
                                   ("50_pct", frame[frame.snapshot_fraction.eq(0.5)]),
                                   ("75_pct_latest_period", frame[frame.snapshot_fraction.eq(0.75)])):
            out["views"][label][stage_label] = {}
            for method in METHODS:
                eligible = stage[np.isfinite(stage[method])]
                out["views"][label][stage_label][method] = measure(eligible.actual, eligible[method]) if len(eligible) else {"n": 0, "status": "NOT_COMPARABLE"}
        comparable = frame[np.isfinite(frame.grade_subject_mean)]
        out["views"][label]["grade_subject_comparable_rf"] = (
            measure(comparable.actual, comparable.rf) if len(comparable) else {"n": 0, "status": "NOT_COMPARABLE"})
    for dimension in ("grade_level", "subject", "school_year", "weight_pattern"):
        out["subgroups"][dimension] = {}
        for value, group in seen.groupby(dimension):
            out["subgroups"][dimension][str(value)] = {"periods": len(group) // 3,
                                                        "snapshots": measure(group.actual, group.rf),
                                                        "latest": measure(group[group.snapshot_fraction.eq(0.75)].actual,
                                                                          group[group.snapshot_fraction.eq(0.75)].rf)}
    for band, mask in (("<75", seen.actual.lt(75)), ("75-79", seen.actual.ge(75) & seen.actual.lt(80)),
                       ("80-84", seen.actual.ge(80) & seen.actual.lt(85)),
                       ("85-89", seen.actual.ge(85) & seen.actual.lt(90)), ("90+", seen.actual.ge(90))):
        frame = seen[mask]
        out["subgroups"].setdefault("actual_grade_band", {})[band] = {
            "periods": len(frame) // 3,
            "snapshots": measure(frame.actual, frame.rf) if len(frame) else {"n": 0, "status": "NOT_EVALUABLE"},
            "latest": measure(frame[frame.snapshot_fraction.eq(0.75)].actual,
                              frame[frame.snapshot_fraction.eq(0.75)].rf) if len(frame) else {"n": 0, "status": "NOT_EVALUABLE"},
        }
    latest = seen[seen.snapshot_fraction.eq(0.75)]
    def band(values):
        return pd.cut(values, [-np.inf, 75, 85, 90, np.inf], right=False,
                      labels=["HIGH_RISK", "MODERATE_RISK", "NEEDS_MONITORING", "LOW_RISK"])
    actual_band, predicted_band = band(latest.actual), band(latest.rf)
    labels = ["HIGH_RISK", "MODERATE_RISK", "NEEDS_MONITORING", "LOW_RISK"]
    matrix = pd.crosstab(actual_band, predicted_band).reindex(index=labels, columns=labels, fill_value=0)
    out["risk"] = {"unit": "latest 75% snapshot, one per student-period",
                   "actual": actual_band.value_counts().reindex(labels, fill_value=0).to_dict(),
                   "predicted": predicted_band.value_counts().reindex(labels, fill_value=0).to_dict(),
                   "confusion_matrix": matrix.to_dict(orient="index"),
                   "agreement_pct": round(float(np.mean(actual_band.to_numpy() == predicted_band.to_numpy()) * 100), 2),
                   "recall_pct": {name: (round(float(matrix.loc[name, name] / matrix.loc[name].sum() * 100), 2)
                                          if matrix.loc[name].sum() else None) for name in labels}}
    return out


def run() -> dict:
    data, verification = verify_data()
    plan = folds(data)
    tables = historical_tables([fold["file"] for fold in plan])
    verification["historical_lookup_parity"] = historical_lookup_parity(data, tables)
    if verification["historical_lookup_parity"]["missing"] or verification["historical_lookup_parity"]["historical_lookup_mismatches"]:
        raise ValueError("Historical workbook lookup does not reproduce cached grades")
    predictions = evaluate(data, plan, tables)
    if {"raw_student_key", "student_period_key", "row_identity", "learner_name", "lrn"}.intersection(predictions.columns):
        raise ValueError("Direct student identifiers entered evaluation output")
    if not np.isfinite(predictions[["rf", "linear", "global_mean"]].to_numpy()).all():
        raise ValueError("Non-finite model or global-baseline prediction")
    verification["fold_membership_sha256"] = hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest()
    verification["direct_pii_in_evaluation_output"] = False
    report = {"model_name": MODEL_NAME, "lifecycle": "DEVELOPMENT", "is_active": False,
              "production_validated": False, "independent_three_term_validation": False,
              "verification": verification, "architecture": {
                  "preprocessing": "median numeric imputation; subject one-hot encode with unknown ignored",
                  "n_estimators": 300, "max_depth": None, "max_features": 0.5,
                  "min_samples_leaf": 2, "min_samples_split": 5, "random_state": RANDOM_STATE,
                  "n_jobs": -1}, **summarize(predictions, plan)}
    # Refit once on all canonical historical snapshots for an isolated development artifact.
    model = rf_pipeline(list(FEATURE_COLUMNS))
    model.fit(data[list(FEATURE_COLUMNS)], data[TARGET_COLUMN])
    artifact = {"model_name": MODEL_NAME, "lifecycle": "DEVELOPMENT", "is_active": False,
                "production_validated": False, "independent_three_term_validation": False,
                "pipeline": model, "target_column": TARGET_COLUMN,
                "feature_columns": list(FEATURE_COLUMNS), "dataset_sha256": verification["data_sha256"],
                "feature_schema_sha256": verification["feature_schema_sha256"], "random_state": RANDOM_STATE}
    MODEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(artifact, MODEL_PATH)
    loaded = joblib.load(MODEL_PATH)
    loaded_pred = loaded["pipeline"].predict(data.iloc[:10][list(FEATURE_COLUMNS)])
    if not np.isfinite(loaded_pred).all() or not np.allclose(loaded_pred, model.predict(data.iloc[:10][list(FEATURE_COLUMNS)])):
        raise ValueError("Saved candidate artifact failed round-trip")
    report["artifact"] = {"path": str(MODEL_PATH), "sha256": sha256(MODEL_PATH),
                          "round_trip_predictions_finite": True}
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    predictions.to_csv(OUTPUT_DIR / "holdout_predictions.csv", index=False, float_format="%.12g")
    (OUTPUT_DIR / "evaluation_report.json").write_text(json.dumps(report, indent=2, sort_keys=True, allow_nan=False) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), indent=2, sort_keys=True))
