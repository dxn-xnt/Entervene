"""Audit historical V3 labels and build an official-workbook-target candidate.

This is an offline data-integrity tool. It never trains or updates a model.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, median
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from app.ml.BuildCurrentTermDevelopmentDataset import FEATURE_COLUMNS
from app.ml.EClassRecordRawParser import EClassRecordRawParser


BACKEND_DIR = Path(__file__).resolve().parents[2]
RAW_DIR = BACKEND_DIR / "data" / "raw_eclass_records"
PROCESSED_SOURCE = BACKEND_DIR / "data" / "processed" / "unified_v3_official_ecr_candidate_with_tle.csv"
DEVELOPMENT_DATASET = BACKEND_DIR / "data" / "datasets" / "current_term_development" / "current_term_development_snapshots.csv"
OUTPUT_DIR = BACKEND_DIR / "data" / "datasets" / "current_term_official_target_candidate"
LINEAGE_FILE = "target_lineage.csv"
CANDIDATE_FILE = "current_term_official_target_candidate_snapshots.csv"

OFFICIAL_SOURCE = "OFFICIAL_WORKBOOK_GRADE"
RECONSTRUCTED_SOURCE = "RECONSTRUCTED_OFFICIAL_20_50_30_TARGETS"
HISTORICAL_OFFICIAL_SOURCE = "HISTORICAL_OFFICIAL_20_60_20_TARGETS"

# The pre-4G backend interpolation anchors in Git commit d2faef8. This is
# retained only to audit old labels; canonical labels never use this function.
LEGACY_ANCHORS = (
    (100, 100), (95, 98), (90, 95), (85, 91), (80, 87),
    (75, 83), (70, 79), (65, 75), (60, 70), (55, 65),
    (50, 60), (45, 55), (40, 50), (35, 45), (30, 40),
    (25, 35), (20, 30), (15, 25), (10, 20), (5, 15), (0, 10),
)

LINEAGE_COLUMNS = (
    "pseudonymous_period_id", "source_workbook", "source_sheet",
    "school_year", "grade_level", "subject", "historical_period",
    "weight_pattern", "official_grade_column", "official_workbook_grade",
    "existing_v3_target", "existing_minus_official", "existing_target_source",
    "existing_target_classification", "legacy_reconstructed_target",
    "legacy_reconstruction_match", "candidate_canonical_target_source",
    "candidate_canonical_target", "resolution_status",
)


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or ()), list(reader)


def _write_csv(path: Path, fields: list[str] | tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n", extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def _grade_text(value: float) -> str:
    return str(int(value)) if value.is_integer() else str(value)


def _key(source_file: str, sheet_name: str, row_number: Any) -> tuple[str, str, int]:
    return source_file, sheet_name, int(row_number)


def _pseudonymous_period_id(key: tuple[str, str, int]) -> str:
    payload = f"v3-official-period-v1\0{key[0]}\0{key[1]}\0{key[2]}"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _legacy_interpolation(initial_grade: float) -> float:
    grade = max(0.0, min(100.0, initial_grade))
    for (high, result_high), (low, result_low) in zip(LEGACY_ANCHORS, LEGACY_ANCHORS[1:]):
        if low <= grade <= high:
            return round(result_low + (grade - low) / (high - low) * (result_high - result_low), 2)
    raise AssertionError("Legacy anchor coverage incomplete")


def _legacy_reconstruction(raw_row: tuple[Any, ...], layout: Any, source_row: dict[str, str], parser: EClassRecordRawParser) -> float | None:
    weighted_initial = 0.0
    for part, weight_column in (("ww", "ww_weight"), ("pt", "pt_weight"), ("qa", "qa_weight")):
        component = layout.components.get(part)
        weight = _number(source_row.get(weight_column))
        if component is None or weight is None:
            return None
        earned, possible, _, _, _ = parser._component_points(raw_row, component, len(component.activity_cols))
        if possible <= 0:
            return None
        weighted_initial += (earned / possible * 100.0) * (weight / 100.0)
    return _legacy_interpolation(round(weighted_initial, 2))


def _band(grade: float) -> str:
    return "<75" if grade < 75 else "75-<85" if grade < 85 else "85-<90" if grade < 90 else ">=90"


def _five_band(grade: float) -> str:
    return "<75" if grade < 75 else "75-79" if grade < 80 else "80-84" if grade < 85 else "85-89" if grade < 90 else "90+"


def _distribution(grades: list[float]) -> dict[str, dict[str, int]]:
    five = Counter(_five_band(grade) for grade in grades)
    risk = Counter(_band(grade) for grade in grades)
    return {
        "five_band": {key: five[key] for key in ("<75", "75-79", "80-84", "85-89", "90+")},
        "risk_band": {key: risk[key] for key in ("<75", "75-<85", "85-<90", ">=90")},
    }


def _check_feature_integrity(old_rows: list[dict[str, str]], candidate_rows: list[dict[str, str]]) -> None:
    forbidden = {"initial_grade", "final_grade", "quarterly_grade", "target_final_period_grade", "source_file", "sheet_name", "row_number", "raw_student_key", "student_period_key", "row_identity"}
    if forbidden.intersection(FEATURE_COLUMNS):
        raise AssertionError("Final outcome, identity, or source provenance entered model features")
    if len(candidate_rows) > len(old_rows):
        raise AssertionError("Candidate gained snapshots")
    old_by_row = {row["row_identity"]: row for row in old_rows}
    student_splits: dict[str, str] = {}
    targets: dict[str, str] = {}
    stages: dict[str, set[str]] = defaultdict(set)
    for row in candidate_rows:
        old = old_by_row[row["row_identity"]]
        for column in FEATURE_COLUMNS:
            if row[column] != old[column]:
                raise AssertionError(f"Feature changed while replacing target: {column}")
        for column in ("raw_student_key", "student_period_key", "source_file", "sheet_name", "row_number", "split", "snapshot_fraction"):
            if row[column] != old[column]:
                raise AssertionError(f"Student grouping or provenance changed: {column}")
        student = row["raw_student_key"]
        if student in student_splits and student_splits[student] != row["split"]:
            raise AssertionError("Student appears in multiple splits")
        student_splits[student] = row["split"]
        period = row["student_period_key"]
        if period in targets and targets[period] != row["target_final_period_grade"]:
            raise AssertionError("Snapshot targets differ within one student-period")
        targets[period] = row["target_final_period_grade"]
        fraction = row["snapshot_fraction"]
        stages[period].add(fraction)
        if fraction in {"0.25", "0.5"}:
            for column in ("qa_available_activity_count", "qa_points_earned_so_far", "qa_points_possible_so_far", "qa_has_evidence"):
                if _number(row[column]) != 0:
                    raise AssertionError("Future QA evidence entered an early snapshot")
    if any(fractions != {"0.25", "0.5", "0.75"} for fractions in stages.values()):
        raise AssertionError("A student-period does not have all three partial snapshots")


def build_candidate(
    *,
    raw_dir: Path = RAW_DIR,
    processed_source: Path = PROCESSED_SOURCE,
    development_dataset: Path = DEVELOPMENT_DATASET,
    output_dir: Path = OUTPUT_DIR,
) -> dict[str, Any]:
    if output_dir.resolve() == development_dataset.parent.resolve():
        raise ValueError("Candidate output must not overwrite the original V3 dataset")
    source_fields, source_rows = _read_csv(processed_source)
    old_fields, old_rows = _read_csv(development_dataset)
    if not old_rows or not source_rows:
        raise ValueError("Processed source and development dataset must be populated")
    if "target_final_period_grade" not in source_fields or "target_final_period_grade" not in old_fields:
        raise ValueError("Target column missing")

    source_by_key: dict[tuple[str, str, int], dict[str, str]] = {}
    for row in source_rows:
        key = _key(row["source_workbook"], row["source_sheet"], row["row_number"])
        previous = source_by_key.setdefault(key, row)
        if previous["target_final_period_grade"] != row["target_final_period_grade"] or previous["target_source"] != row["target_source"]:
            raise AssertionError("Processed source has inconsistent targets within a student-period")

    periods: dict[tuple[str, str, int], dict[str, str]] = {}
    snapshot_count_by_key: Counter[tuple[str, str, int]] = Counter()
    for row in old_rows:
        key = _key(row["source_file"], row["sheet_name"], row["row_number"])
        previous = periods.setdefault(key, row)
        if previous["student_period_key"] != row["student_period_key"] or previous["target_final_period_grade"] != row["target_final_period_grade"]:
            raise AssertionError("Development snapshots have inconsistent period identity or target")
        source = source_by_key.get(key)
        if source is None or _number(source["target_final_period_grade"]) != _number(row["target_final_period_grade"]):
            raise AssertionError("Development target cannot be traced to the processed source")
        snapshot_count_by_key[key] += 1

    parser = EClassRecordRawParser()
    extracted: dict[tuple[str, str, int], tuple[float | None, str, float | None, str]] = {}
    by_workbook: dict[str, dict[str, list[tuple[str, str, int]]]] = defaultdict(lambda: defaultdict(list))
    for key in periods:
        by_workbook[key[0]][key[1]].append(key)
    for filename, sheets in sorted(by_workbook.items()):
        workbook_path = raw_dir / filename
        if not workbook_path.is_file():
            for keys in sheets.values():
                for key in keys:
                    extracted[key] = (None, "", None, "UNRESOLVED_WORKBOOK_MISSING")
            continue
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            for sheet_name, keys in sorted(sheets.items()):
                if sheet_name not in workbook:
                    for key in keys:
                        extracted[key] = (None, "", None, "UNRESOLVED_SHEET_MISSING")
                    continue
                all_rows = list(workbook[sheet_name].iter_rows(values_only=True))
                try:
                    layout = parser._find_layout(all_rows)
                except ValueError:
                    for key in keys:
                        extracted[key] = (None, "", None, "UNRESOLVED_GRADE_LAYOUT")
                    continue
                column = get_column_letter(layout.final_grade_col + 1)
                header = str(all_rows[layout.label_row_idx][layout.final_grade_col] or "").upper()
                if "GRADE" not in header:
                    raise AssertionError("Located historical official-grade column lacks a grade heading")
                for key in keys:
                    if key[2] < 1 or key[2] > len(all_rows):
                        extracted[key] = (None, column, None, "UNRESOLVED_ROW_MISSING")
                        continue
                    raw_row = all_rows[key[2] - 1]
                    value = raw_row[layout.final_grade_col] if layout.final_grade_col < len(raw_row) else None
                    official = _number(value)
                    status = "RESOLVED" if official is not None and 0 <= official <= 100 else "UNRESOLVED_OFFICIAL_GRADE"
                    legacy = None
                    if periods[key]["target_source"] == RECONSTRUCTED_SOURCE:
                        legacy = _legacy_reconstruction(raw_row, layout, source_by_key[key], parser)
                    extracted[key] = (official if status == "RESOLVED" else None, column, legacy, status)
        finally:
            workbook.close()

    lineage: list[dict[str, Any]] = []
    canonical_by_key: dict[tuple[str, str, int], float] = {}
    for key, old in sorted(periods.items()):
        official, column, legacy, status = extracted[key]
        old_target = _number(old["target_final_period_grade"])
        if old_target is None:
            raise AssertionError("Existing V3 target is nonnumeric")
        existing_source = old["target_source"]
        legacy_match = legacy is not None and legacy == old_target
        if existing_source == RECONSTRUCTED_SOURCE:
            classification = "RECONSTRUCTED_TARGET" if legacy_match else "TARGET_SOURCE_UNRESOLVED"
        elif existing_source == HISTORICAL_OFFICIAL_SOURCE:
            classification = "OFFICIAL_WORKBOOK_GRADE" if official == old_target else "TARGET_SOURCE_UNRESOLVED"
        else:
            classification = "TARGET_SOURCE_UNRESOLVED"
        if official is not None:
            canonical_by_key[key] = official
        lineage.append({
            "pseudonymous_period_id": _pseudonymous_period_id(key),
            "source_workbook": key[0], "source_sheet": key[1],
            "school_year": old["school_year"], "grade_level": old["grade_level"],
            "subject": old["subject"], "historical_period": old["period_sequence_audit"],
            "weight_pattern": old["weight_pattern"], "official_grade_column": column,
            "official_workbook_grade": "" if official is None else _grade_text(official),
            "existing_v3_target": old["target_final_period_grade"],
            "existing_minus_official": "" if official is None else str(round(old_target - official, 4)),
            "existing_target_source": existing_source,
            "existing_target_classification": classification,
            "legacy_reconstructed_target": "" if legacy is None else _grade_text(legacy),
            "legacy_reconstruction_match": "" if legacy is None else str(legacy_match).lower(),
            "candidate_canonical_target_source": OFFICIAL_SOURCE if official is not None else "",
            "candidate_canonical_target": "" if official is None else _grade_text(official),
            "resolution_status": status,
        })

    candidate_rows: list[dict[str, str]] = []
    for old in old_rows:
        key = _key(old["source_file"], old["sheet_name"], old["row_number"])
        official = canonical_by_key.get(key)
        if official is None:
            continue
        row = dict(old)
        row["target_source"] = OFFICIAL_SOURCE
        row["target_final_period_grade"] = _grade_text(official)
        candidate_rows.append(row)
    _check_feature_integrity(old_rows, candidate_rows)

    output_dir.mkdir(parents=True, exist_ok=True)
    lineage_path = output_dir / LINEAGE_FILE
    candidate_path = output_dir / CANDIDATE_FILE
    _write_csv(lineage_path, LINEAGE_COLUMNS, lineage)
    _write_csv(candidate_path, old_fields, candidate_rows)

    resolved = [entry for entry in lineage if entry["resolution_status"] == "RESOLVED"]
    old_grades = [float(entry["existing_v3_target"]) for entry in resolved]
    new_grades = [float(entry["candidate_canonical_target"]) for entry in resolved]
    changed = [entry for entry in resolved if float(entry["existing_v3_target"]) != float(entry["candidate_canonical_target"])]
    absolute_differences = [abs(old - new) for old, new in zip(old_grades, new_grades)]
    by_file: dict[str, dict[str, Any]] = {}
    for filename in sorted(by_workbook):
        entries = [entry for entry in lineage if entry["source_workbook"] == filename]
        by_file[filename] = {
            "student_periods": len(entries),
            "official_resolved": sum(entry["resolution_status"] == "RESOLVED" for entry in entries),
            "target_matches_official": sum(entry["resolution_status"] == "RESOLVED" and float(entry["existing_v3_target"]) == float(entry["candidate_canonical_target"]) for entry in entries),
            "target_differs_official": sum(entry["resolution_status"] == "RESOLVED" and float(entry["existing_v3_target"]) != float(entry["candidate_canonical_target"]) for entry in entries),
            "grade_columns": sorted({entry["official_grade_column"] for entry in entries if entry["official_grade_column"]}),
            "school_years": sorted({entry["school_year"] for entry in entries}),
            "grades": sorted({entry["grade_level"] for entry in entries}),
            "subjects": sorted({entry["subject"] for entry in entries}),
            "weight_patterns": sorted({entry["weight_pattern"] for entry in entries}),
            "historical_periods": sorted({entry["historical_period"] for entry in entries}),
            "existing_target_classifications": dict(sorted(Counter(entry["existing_target_classification"] for entry in entries).items())),
        }
    summary = {
        "source_snapshots": len(source_rows), "development_snapshots": len(old_rows),
        "student_periods": len(periods), "resolved_student_periods": len(resolved),
        "unresolved_student_periods": len(periods) - len(resolved),
        "candidate_snapshots": len(candidate_rows),
        "unchanged_student_periods": len(resolved) - len(changed),
        "changed_student_periods": len(changed),
        "unchanged_snapshots": sum(snapshot_count_by_key[key] for key, entry in zip(sorted(periods), lineage) if entry["resolution_status"] == "RESOLVED" and float(entry["existing_v3_target"]) == float(entry["candidate_canonical_target"])),
        "changed_snapshots": sum(snapshot_count_by_key[key] for key, entry in zip(sorted(periods), lineage) if entry["resolution_status"] == "RESOLVED" and float(entry["existing_v3_target"]) != float(entry["candidate_canonical_target"])),
        "old_new_label_mae": mean(absolute_differences) if absolute_differences else None,
        "old_new_label_median_absolute": median(absolute_differences) if absolute_differences else None,
        "old_new_label_max_absolute": max(absolute_differences) if absolute_differences else None,
        "old_minus_new_mean_signed": mean(old - new for old, new in zip(old_grades, new_grades)) if old_grades else None,
        "risk_band_changed_student_periods": sum(_band(old) != _band(new) for old, new in zip(old_grades, new_grades)),
        "risk_band_changed_snapshots": sum(snapshot_count_by_key[key] for key, entry in zip(sorted(periods), lineage) if entry["resolution_status"] == "RESOLVED" and _band(float(entry["existing_v3_target"])) != _band(float(entry["candidate_canonical_target"]))),
        "old_target_distribution_student_periods": _distribution(old_grades),
        "canonical_target_distribution_student_periods": _distribution(new_grades),
        "canonical_target_distribution_snapshots": _distribution([float(row["target_final_period_grade"]) for row in candidate_rows]),
        "weight_patterns": dict(sorted(Counter(entry["weight_pattern"] for entry in lineage).items())),
        "legacy_reconstruction_matches": sum(entry["legacy_reconstruction_match"] == "true" for entry in lineage),
        "legacy_reconstruction_total": sum(entry["existing_target_source"] == RECONSTRUCTED_SOURCE for entry in lineage),
        "changed_by_subject": dict(sorted(Counter(entry["subject"] for entry in changed).items())),
        "periods_by_grade": dict(sorted(Counter(entry["grade_level"] for entry in resolved).items())),
        "periods_by_subject": dict(sorted(Counter(entry["subject"] for entry in resolved).items())),
        "periods_by_school_year": dict(sorted(Counter(entry["school_year"] for entry in resolved).items())),
        "snapshots_by_grade": dict(sorted(Counter(row["grade_level"] for row in candidate_rows).items())),
        "snapshots_by_subject": dict(sorted(Counter(row["subject"] for row in candidate_rows).items())),
        "snapshots_by_school_year": dict(sorted(Counter(row["school_year"] for row in candidate_rows).items())),
        "by_workbook": by_file,
        "lineage_sha256": hashlib.sha256(lineage_path.read_bytes()).hexdigest(),
        "candidate_sha256": hashlib.sha256(candidate_path.read_bytes()).hexdigest(),
    }
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=RAW_DIR)
    parser.add_argument("--processed-source", type=Path, default=PROCESSED_SOURCE)
    parser.add_argument("--development-dataset", type=Path, default=DEVELOPMENT_DATASET)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    summary = build_candidate(raw_dir=args.raw_dir, processed_source=args.processed_source, development_dataset=args.development_dataset, output_dir=args.output_dir)
    print(json.dumps(summary, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
