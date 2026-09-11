from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.models.academic.Class_ import Class
from app.models.academic.GradingTemplate import GradingTemplate
from app.models.academic.GradingTemplateComponent import GradingTemplateComponent
from app.models.academic.Subject import Subject
from app.models.academic.SubjectLoad import SubjectLoad
from app.models.people.AcademicStaff import AcademicStaff
from app.models.settings.Setting import Setting
from app.services.classes.ClassQueryService import _student_gender_group
from app.services.student_record.StudentRecordService import (
    _deped_transmuted,
    _match_component_category,
    get_performance_descriptor,
    resolve_subject_grading_weights,
    teacher_student_gradebook,
)


# ── DepEd Styling Constants ──────────────────────────────────────────

FONT_FAMILY = "Calibri"

FONT_TITLE = Font(name=FONT_FAMILY, size=14, bold=True, color="1F497D")
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=11, bold=True, color="1F497D")
FONT_HEADER_NOTE = Font(name=FONT_FAMILY, size=9, italic=True, color="595959")
FONT_META_LABEL = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
FONT_META_VAL = Font(name=FONT_FAMILY, size=10, bold=False, color="000000")

FONT_TABLE_HEADER = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
FONT_HPS = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
FONT_GENDER_SECTION = Font(name=FONT_FAMILY, size=10, bold=True, color="1F497D")
FONT_ROW_TEXT = Font(name=FONT_FAMILY, size=10, bold=False, color="000000")
FONT_ROW_NUM = Font(name=FONT_FAMILY, size=10, bold=False, color="000000")
FONT_SUMMARY = Font(name=FONT_FAMILY, size=10, bold=True, color="333333")

# DepEd Component Soft Colors
FILL_BLUE = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")       # Written Works
FILL_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")      # Performance Tasks
FILL_PEACH = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")      # Quarterly Assessment
FILL_GOLD = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")       # Final Grades / Summary
FILL_GRAY = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")       # HPS / Fallback
FILL_GENDER = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")     # Male/Female dividers

# Cell Borders
BORDER_THIN_BLACK = Side(border_style="thin", color="000000")
BORDER_MEDIUM_BLACK = Side(border_style="medium", color="000000")
BORDER_DOUBLE_BLACK = Side(border_style="double", color="000000")

CELL_BORDER = Border(
    left=BORDER_THIN_BLACK,
    right=BORDER_THIN_BLACK,
    top=BORDER_THIN_BLACK,
    bottom=BORDER_THIN_BLACK,
)

HPS_BORDER = Border(
    left=BORDER_THIN_BLACK,
    right=BORDER_THIN_BLACK,
    top=BORDER_THIN_BLACK,
    bottom=BORDER_MEDIUM_BLACK,
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_NAME = Alignment(horizontal="left", vertical="center", indent=1)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


@dataclass
class ExportComponentColGroup:
    name: str
    weight_pct: float
    category: str
    assignments: list[Any]
    sub_col_count: int
    start_col: int
    end_col: int
    total_col: int
    ps_col: int
    ws_col: int
    fill: PatternFill


def _style_range(
    ws: Worksheet,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    font: Font | None = None,
    fill: PatternFill | None = None,
    border: Border | None = None,
    alignment: Alignment | None = None,
) -> None:
    """Apply styling across an entire range (including merged areas) to prevent missing borders."""
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if alignment is not None:
                cell.alignment = alignment


def _get_school_metadata(db: Session, scope: Any) -> dict[str, str]:
    """
    Retrieve real school settings if present.
    If not configured in Setting table, return empty strings ("") rather than
    fabricated/false regional placeholders.
    """
    settings_rows = {s.key: (s.value or "").strip() for s in db.query(Setting).all()}

    # School Name: setting school_name -> app_name -> fallback empty
    school_name = settings_rows.get("school_name") or settings_rows.get("app_name") or ""

    # Region & Division: Only use if actually configured in settings
    region = settings_rows.get("school_region") or settings_rows.get("region") or settings_rows.get("deped_region") or ""
    division = settings_rows.get("school_division") or settings_rows.get("division") or settings_rows.get("deped_division") or ""
    school_id = settings_rows.get("school_id") or ""

    # School Year
    school_year = (getattr(scope, "year_label", None) or settings_rows.get("current_school_year") or "").strip()

    # Section, Subject, Period
    section_name = (getattr(scope, "section_name", None) or "").strip()
    subject_name = (getattr(scope, "subject_name", None) or "").strip()
    period_name = (getattr(scope, "period_name", None) or "").strip()

    # Teacher Name
    teacher_name = ""
    if hasattr(scope, "original_teacher_name") and scope.original_teacher_name:
        teacher_name = scope.original_teacher_name
    elif hasattr(scope, "acting_staff_id") and scope.acting_staff_id:
        staff = db.get(AcademicStaff, scope.acting_staff_id)
        if staff:
            teacher_name = f"{staff.first_name} {staff.last_name}".strip()

    return {
        "school_name": school_name,
        "region": region,
        "division": division,
        "school_id": school_id,
        "school_year": school_year,
        "section_name": section_name,
        "subject_name": subject_name,
        "period_name": period_name,
        "teacher_name": teacher_name,
    }


def generate_class_record_sheet(
    wb: openpyxl.Workbook,
    db: Session,
    class_id: int,
    subject_id: int,
    academic_period_id: int,
    staff_id: str | None = None,
    sheet_title: str | None = None,
) -> Worksheet:
    """
    Generate an official DepEd single-term Class Record worksheet.
    
    This function is reusable and parameterized so that it can be called
    repeatedly for multi-term workbooks without modification.
    """
    # 1. Resolve teacher staff_id if not provided or admin caller
    effective_staff_id = staff_id
    if not effective_staff_id:
        load = (
            db.query(SubjectLoad)
            .filter(
                SubjectLoad.class_id == class_id,
                SubjectLoad.subject_id == subject_id,
                SubjectLoad.academic_period_id == academic_period_id,
                SubjectLoad.status.in_(["active", "published"]),
            )
            .first()
        )
        if load and load.staff_id:
            effective_staff_id = load.staff_id

    # 2. Fetch gradebook data from StudentRecordService
    gradebook = teacher_student_gradebook(
        db=db,
        staff_id=effective_staff_id or "",
        class_id=class_id,
        subject_id=subject_id,
        academic_period_id=academic_period_id,
    )

    scope = gradebook.scope
    meta = _get_school_metadata(db, scope)

    # 3. Create or reuse worksheet
    term_title = sheet_title or meta["period_name"] or "Term"
    clean_sheet_title = re.sub(r"[\\/*?:\[\]]", "", term_title)[:31].strip() or "Class Record"

    if len(wb.sheetnames) == 1 and wb.sheetnames[0] == "Sheet":
        ws = wb.active
        ws.title = clean_sheet_title
    else:
        ws = wb.create_sheet(title=clean_sheet_title)

    # Ensure Excel gridlines are displayed
    ws.views.sheetView[0].showGridLines = True

    # 4. Resolve dynamic GradingTemplate components
    class_obj = db.get(Class, class_id)
    academic_level_id = getattr(class_obj, "academic_level_id", None)
    weights = resolve_subject_grading_weights(db, subject_id, academic_level_id)

    db_components: list[GradingTemplateComponent] = []
    if weights.template_id:
        db_components = (
            db.query(GradingTemplateComponent)
            .filter(GradingTemplateComponent.grading_template_id == weights.template_id)
            .order_by(GradingTemplateComponent.display_order)
            .all()
        )

    # Fallback to standard 3 components if no template records found
    if not db_components:
        component_defs = [
            ("Written Works", weights.ww_weight * 100.0, "WW"),
            ("Performance Tasks", weights.pt_weight * 100.0, "PT"),
            ("Quarterly Assessment", weights.qa_weight * 100.0, "QA"),
        ]
    else:
        component_defs = [
            (c.component_name, float(c.weight), _match_component_category(c.component_name) or "WW")
            for c in db_components
        ]

    # Category header assignments from gradebook
    cg = gradebook.classwork[0] if gradebook.classwork else None
    ww_assignments = cg.writtenWork if cg else []
    pt_assignments = cg.performanceTask if cg else []
    qa_assignments = (cg.quarterlyAssessment or cg.exams) if cg else []

    # Map color palette to components
    palette = [FILL_BLUE, FILL_GREEN, FILL_PEACH, FILL_GRAY]

    # 5. Build dynamic component column groups (Option A: adaptive count, minimum 1)
    comp_groups: list[ExportComponentColGroup] = []
    current_col = 3  # Column 1 = No., Column 2 = Learners' Names

    for idx, (c_name, c_weight, c_cat) in enumerate(component_defs):
        if c_cat == "WW":
            asgns = ww_assignments
        elif c_cat == "PT":
            asgns = pt_assignments
        elif c_cat == "QA":
            asgns = qa_assignments
        else:
            asgns = []

        sub_count = max(1, len(asgns))
        start_col = current_col
        total_col = current_col + sub_count
        ps_col = total_col + 1
        ws_col = ps_col + 1
        end_col = ws_col

        fill = palette[idx % len(palette)]

        comp_groups.append(
            ExportComponentColGroup(
                name=c_name,
                weight_pct=c_weight,
                category=c_cat,
                assignments=asgns,
                sub_col_count=sub_count,
                start_col=start_col,
                end_col=end_col,
                total_col=total_col,
                ps_col=ps_col,
                ws_col=ws_col,
                fill=fill,
            )
        )
        current_col = end_col + 1

    initial_grade_col = current_col
    term_grade_col = current_col + 1
    descriptor_col = current_col + 2
    total_cols = descriptor_col

    # ── 6. Render Header Block (Rows 1 to 8) ──────────────────────────
    ws.cell(1, 1).value = "Republic of the Philippines"
    ws.cell(1, 1).font = FONT_HEADER_NOTE

    ws.cell(2, 1).value = "Department of Education"
    ws.cell(2, 1).font = FONT_SUBTITLE

    ws.cell(3, 1).value = "ELECTRONIC CLASS RECORD"
    ws.cell(3, 1).font = FONT_TITLE

    # Metadata Grid (Rows 5 to 8)
    # Left Column: Region, Division, School Name, School ID
    ws.cell(5, 1).value = "REGION:"
    ws.cell(5, 1).font = FONT_META_LABEL
    ws.cell(5, 2).value = meta["region"]
    ws.cell(5, 2).font = FONT_META_VAL

    ws.cell(6, 1).value = "DIVISION:"
    ws.cell(6, 1).font = FONT_META_LABEL
    ws.cell(6, 2).value = meta["division"]
    ws.cell(6, 2).font = FONT_META_VAL

    ws.cell(7, 1).value = "SCHOOL NAME:"
    ws.cell(7, 1).font = FONT_META_LABEL
    ws.cell(7, 2).value = meta["school_name"]
    ws.cell(7, 2).font = FONT_META_VAL

    ws.cell(8, 1).value = "SCHOOL ID:"
    ws.cell(8, 1).font = FONT_META_LABEL
    ws.cell(8, 2).value = meta["school_id"]
    ws.cell(8, 2).font = FONT_META_VAL

    # Right Column: School Year, Grade & Section, Teacher, Subject & Quarter
    mid_label_col = max(5, min(8, total_cols - 4))
    mid_val_col = mid_label_col + 1

    ws.cell(5, mid_label_col).value = "SCHOOL YEAR:"
    ws.cell(5, mid_label_col).font = FONT_META_LABEL
    ws.cell(5, mid_val_col).value = meta["school_year"]
    ws.cell(5, mid_val_col).font = FONT_META_VAL

    ws.cell(6, mid_label_col).value = "GRADE & SECTION:"
    ws.cell(6, mid_label_col).font = FONT_META_LABEL
    ws.cell(6, mid_val_col).value = meta["section_name"]
    ws.cell(6, mid_val_col).font = FONT_META_VAL

    ws.cell(7, mid_label_col).value = "TEACHER:"
    ws.cell(7, mid_label_col).font = FONT_META_LABEL
    ws.cell(7, mid_val_col).value = meta["teacher_name"]
    ws.cell(7, mid_val_col).font = FONT_META_VAL

    ws.cell(8, mid_label_col).value = "SUBJECT & QUARTER:"
    ws.cell(8, mid_label_col).font = FONT_META_LABEL
    ws.cell(8, mid_val_col).value = f"{meta['subject_name']} - {meta['period_name']}".strip(" -")
    ws.cell(8, mid_val_col).font = FONT_META_VAL

    # ── 7. Render Table Headers (Rows 10 & 11) ────────────────────────
    # Row 10: Component Group Headers & Merged Columns
    # Row 11: Sub-columns (1..k, Total, PS, WS)

    # Col 1: NO.
    ws.merge_cells(start_row=10, start_column=1, end_row=11, end_column=1)
    ws.cell(10, 1).value = "NO."
    _style_range(ws, 1, 10, 1, 11, font=FONT_TABLE_HEADER, fill=FILL_BLUE, border=CELL_BORDER, alignment=ALIGN_CENTER)

    # Col 2: LEARNERS' NAMES
    ws.merge_cells(start_row=10, start_column=2, end_row=11, end_column=2)
    ws.cell(10, 2).value = "LEARNERS' NAMES"
    _style_range(ws, 2, 10, 2, 11, font=FONT_TABLE_HEADER, fill=FILL_BLUE, border=CELL_BORDER, alignment=ALIGN_CENTER)

    # Dynamic Components
    for comp in comp_groups:
        # Merged header in Row 10 across all its sub-columns
        ws.merge_cells(start_row=10, start_column=comp.start_col, end_row=10, end_column=comp.end_col)
        weight_str = f"{int(round(comp.weight_pct))}%" if comp.weight_pct.is_integer() else f"{comp.weight_pct}%"
        ws.cell(10, comp.start_col).value = f"{comp.name.upper()} ({weight_str})"
        _style_range(
            ws,
            comp.start_col,
            10,
            comp.end_col,
            10,
            font=FONT_TABLE_HEADER,
            fill=comp.fill,
            border=CELL_BORDER,
            alignment=ALIGN_CENTER,
        )

        # Row 11: Sub-column headers
        for j in range(comp.sub_col_count):
            sub_col = comp.start_col + j
            ws.cell(11, sub_col).value = j + 1
            _style_range(ws, sub_col, 11, sub_col, 11, font=FONT_TABLE_HEADER, fill=comp.fill, border=CELL_BORDER, alignment=ALIGN_CENTER)

        ws.cell(11, comp.total_col).value = "Total"
        _style_range(ws, comp.total_col, 11, comp.total_col, 11, font=FONT_TABLE_HEADER, fill=comp.fill, border=CELL_BORDER, alignment=ALIGN_CENTER)

        ws.cell(11, comp.ps_col).value = "PS"
        _style_range(ws, comp.ps_col, 11, comp.ps_col, 11, font=FONT_TABLE_HEADER, fill=comp.fill, border=CELL_BORDER, alignment=ALIGN_CENTER)

        ws.cell(11, comp.ws_col).value = "WS"
        _style_range(ws, comp.ws_col, 11, comp.ws_col, 11, font=FONT_TABLE_HEADER, fill=comp.fill, border=CELL_BORDER, alignment=ALIGN_CENTER)

    # Summary Columns (Merged Rows 10 & 11)
    ws.merge_cells(start_row=10, start_column=initial_grade_col, end_row=11, end_column=initial_grade_col)
    ws.cell(10, initial_grade_col).value = "INITIAL\nGRADE"
    _style_range(ws, initial_grade_col, 10, initial_grade_col, 11, font=FONT_TABLE_HEADER, fill=FILL_GOLD, border=CELL_BORDER, alignment=ALIGN_CENTER)

    ws.merge_cells(start_row=10, start_column=term_grade_col, end_row=11, end_column=term_grade_col)
    ws.cell(10, term_grade_col).value = "QUARTERLY\nGRADE"
    _style_range(ws, term_grade_col, 10, term_grade_col, 11, font=FONT_TABLE_HEADER, fill=FILL_GOLD, border=CELL_BORDER, alignment=ALIGN_CENTER)

    ws.merge_cells(start_row=10, start_column=descriptor_col, end_row=11, end_column=descriptor_col)
    ws.cell(10, descriptor_col).value = "DESCRIPTOR"
    _style_range(ws, descriptor_col, 10, descriptor_col, 11, font=FONT_TABLE_HEADER, fill=FILL_GOLD, border=CELL_BORDER, alignment=ALIGN_CENTER)

    # ── 8. Render Highest Possible Score (HPS) Row (Row 12) ───────────
    hps_row = 12
    ws.cell(hps_row, 1).value = ""
    ws.cell(hps_row, 2).value = "HIGHEST POSSIBLE SCORE"
    _style_range(ws, 1, hps_row, 2, hps_row, font=FONT_HPS, fill=FILL_GRAY, border=HPS_BORDER, alignment=ALIGN_LEFT)

    for comp in comp_groups:
        comp_total_max = 0.0
        for j in range(comp.sub_col_count):
            sub_col = comp.start_col + j
            if j < len(comp.assignments):
                pts = float(comp.assignments[j].maxScore)
                ws.cell(hps_row, sub_col).value = pts
                comp_total_max += pts
            else:
                ws.cell(hps_row, sub_col).value = ""
            _style_range(ws, sub_col, hps_row, sub_col, hps_row, font=FONT_HPS, fill=FILL_GRAY, border=HPS_BORDER, alignment=ALIGN_CENTER)

        # Total
        ws.cell(hps_row, comp.total_col).value = comp_total_max if comp.assignments else ""
        _style_range(ws, comp.total_col, hps_row, comp.total_col, hps_row, font=FONT_HPS, fill=FILL_GRAY, border=HPS_BORDER, alignment=ALIGN_CENTER)

        # PS
        ws.cell(hps_row, comp.ps_col).value = 100.0 if comp.assignments else ""
        _style_range(ws, comp.ps_col, hps_row, comp.ps_col, hps_row, font=FONT_HPS, fill=FILL_GRAY, border=HPS_BORDER, alignment=ALIGN_CENTER)

        # WS
        ws.cell(hps_row, comp.ws_col).value = comp.weight_pct if comp.assignments else ""
        _style_range(ws, comp.ws_col, hps_row, comp.ws_col, hps_row, font=FONT_HPS, fill=FILL_GRAY, border=HPS_BORDER, alignment=ALIGN_CENTER)

    # HPS Summary Columns
    ws.cell(hps_row, initial_grade_col).value = 100.0
    _style_range(ws, initial_grade_col, hps_row, initial_grade_col, hps_row, font=FONT_HPS, fill=FILL_GOLD, border=HPS_BORDER, alignment=ALIGN_CENTER)

    ws.cell(hps_row, term_grade_col).value = 100
    _style_range(ws, term_grade_col, hps_row, term_grade_col, hps_row, font=FONT_HPS, fill=FILL_GOLD, border=HPS_BORDER, alignment=ALIGN_CENTER)

    ws.cell(hps_row, descriptor_col).value = ""
    _style_range(ws, descriptor_col, hps_row, descriptor_col, hps_row, font=FONT_HPS, fill=FILL_GOLD, border=HPS_BORDER, alignment=ALIGN_CENTER)

    # ── 9. Render Student Rows Grouped by Gender ──────────────────────
    males = []
    females = []
    for sg in gradebook.studentGrades:
        grp = _student_gender_group(sg.gender)
        if grp == "Female":
            females.append(sg)
        else:
            males.append(sg)

    males.sort(key=lambda s: s.name.casefold())
    females.sort(key=lambda s: s.name.casefold())

    current_row = 13

    def _render_gender_section(label: str, students: list[Any]) -> None:
        nonlocal current_row
        # Gender Header Row
        ws.cell(current_row, 1).value = ""
        ws.cell(current_row, 2).value = label
        _style_range(ws, 1, current_row, total_cols, current_row, font=FONT_GENDER_SECTION, fill=FILL_GENDER, border=CELL_BORDER, alignment=ALIGN_LEFT)
        current_row += 1

        for idx, student in enumerate(students):
            row_num = current_row
            ws.cell(row_num, 1).value = idx + 1
            ws.cell(row_num, 1).alignment = ALIGN_CENTER
            ws.cell(row_num, 1).font = FONT_ROW_NUM
            ws.cell(row_num, 1).border = CELL_BORDER

            ws.cell(row_num, 2).value = student.name
            ws.cell(row_num, 2).alignment = ALIGN_NAME
            ws.cell(row_num, 2).font = FONT_ROW_TEXT
            ws.cell(row_num, 2).border = CELL_BORDER

            student_ws_sum = 0.0
            has_any_score = False

            for comp in comp_groups:
                if comp.category == "WW":
                    raw_scores = student.writtenWork
                    student_ps = student.ps_written
                elif comp.category == "PT":
                    raw_scores = student.performanceTask
                    student_ps = student.ps_performance
                elif comp.category == "QA":
                    raw_scores = student.quarterlyAssessment or student.exams
                    student_ps = student.ps_quarterly
                else:
                    raw_scores = []
                    student_ps = None

                earned_total = 0.0
                comp_has_score = False

                for j in range(comp.sub_col_count):
                    sub_col = comp.start_col + j
                    score_val = raw_scores[j] if j < len(raw_scores) else None
                    if score_val is not None:
                        ws.cell(row_num, sub_col).value = float(score_val)
                        earned_total += float(score_val)
                        comp_has_score = True
                        has_any_score = True
                    else:
                        ws.cell(row_num, sub_col).value = ""
                    ws.cell(row_num, sub_col).alignment = ALIGN_CENTER
                    ws.cell(row_num, sub_col).font = FONT_ROW_TEXT
                    ws.cell(row_num, sub_col).border = CELL_BORDER

                # Component Total
                total_val = earned_total if comp_has_score else ""
                ws.cell(row_num, comp.total_col).value = total_val
                ws.cell(row_num, comp.total_col).alignment = ALIGN_CENTER
                ws.cell(row_num, comp.total_col).font = FONT_ROW_TEXT
                ws.cell(row_num, comp.total_col).border = CELL_BORDER
                if isinstance(total_val, float):
                    ws.cell(row_num, comp.total_col).number_format = "0.00"

                # Component PS
                if student_ps is not None:
                    ps_val = float(student_ps)
                elif comp_has_score and comp.assignments:
                    hps_tot = sum(float(a.maxScore) for a in comp.assignments)
                    ps_val = round((earned_total / hps_tot) * 100.0, 2) if hps_tot > 0 else 0.0
                else:
                    ps_val = None

                ws.cell(row_num, comp.ps_col).value = ps_val if ps_val is not None else ""
                ws.cell(row_num, comp.ps_col).alignment = ALIGN_CENTER
                ws.cell(row_num, comp.ps_col).font = FONT_ROW_TEXT
                ws.cell(row_num, comp.ps_col).border = CELL_BORDER
                if ps_val is not None:
                    ws.cell(row_num, comp.ps_col).number_format = "0.00"

                # Component WS
                if ps_val is not None:
                    comp_ws = round(ps_val * (comp.weight_pct / 100.0), 2)
                    student_ws_sum += comp_ws
                else:
                    comp_ws = None

                ws.cell(row_num, comp.ws_col).value = comp_ws if comp_ws is not None else ""
                ws.cell(row_num, comp.ws_col).alignment = ALIGN_CENTER
                ws.cell(row_num, comp.ws_col).font = FONT_ROW_TEXT
                ws.cell(row_num, comp.ws_col).border = CELL_BORDER
                if comp_ws is not None:
                    ws.cell(row_num, comp.ws_col).number_format = "0.00"

            # Initial Grade
            ig_val = student.initial_grade if student.initial_grade is not None else (round(student_ws_sum, 2) if has_any_score else None)
            ws.cell(row_num, initial_grade_col).value = ig_val if ig_val is not None else ""
            ws.cell(row_num, initial_grade_col).alignment = ALIGN_CENTER
            ws.cell(row_num, initial_grade_col).font = FONT_ROW_NUM
            ws.cell(row_num, initial_grade_col).border = CELL_BORDER
            if ig_val is not None:
                ws.cell(row_num, initial_grade_col).number_format = "0.00"

            # Quarterly / Term Grade
            tg_val = student.transmuted_grade if student.transmuted_grade is not None else (_deped_transmuted(ig_val) if ig_val is not None else None)
            ws.cell(row_num, term_grade_col).value = tg_val if tg_val is not None else ""
            ws.cell(row_num, term_grade_col).alignment = ALIGN_CENTER
            ws.cell(row_num, term_grade_col).font = FONT_ROW_NUM
            ws.cell(row_num, term_grade_col).border = CELL_BORDER
            if tg_val is not None:
                ws.cell(row_num, term_grade_col).number_format = "0"

            # Performance Descriptor
            desc_val = student.performance_descriptor or (get_performance_descriptor(tg_val) if tg_val is not None else "")
            ws.cell(row_num, descriptor_col).value = desc_val or ""
            ws.cell(row_num, descriptor_col).alignment = ALIGN_CENTER
            ws.cell(row_num, descriptor_col).font = FONT_ROW_TEXT
            ws.cell(row_num, descriptor_col).border = CELL_BORDER

            current_row += 1

    # Render MALE group
    _render_gender_section("MALE", males)

    # Render FEMALE group
    _render_gender_section("FEMALE", females)

    # Summary Footer Row
    ws.cell(current_row, 2).value = (
        f"TOTAL MALE: {len(males)}   |   TOTAL FEMALE: {len(females)}   |   "
        f"TOTAL LEARNERS: {len(males) + len(females)}"
    )
    _style_range(ws, 1, current_row, total_cols, current_row, font=FONT_SUMMARY, fill=FILL_GRAY, border=CELL_BORDER, alignment=ALIGN_LEFT)

    # ── 10. Column Width Optimization ─────────────────────────────────
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions[get_column_letter(initial_grade_col)].width = 14
    ws.column_dimensions[get_column_letter(term_grade_col)].width = 15
    ws.column_dimensions[get_column_letter(descriptor_col)].width = 18

    for comp in comp_groups:
        for j in range(comp.sub_col_count):
            ws.column_dimensions[get_column_letter(comp.start_col + j)].width = 8
        ws.column_dimensions[get_column_letter(comp.total_col)].width = 9
        ws.column_dimensions[get_column_letter(comp.ps_col)].width = 8
        ws.column_dimensions[get_column_letter(comp.ws_col)].width = 8

    return ws


def export_class_record_single_term(
    db: Session,
    class_id: int,
    subject_id: int,
    academic_period_id: int,
    staff_id: str,
) -> tuple[io.BytesIO, str]:
    """
    Generate an in-memory .xlsx file for a single-term DepEd Class Record.
    Returns (BytesIO stream, suggested filename).
    """
    wb = openpyxl.Workbook()

    ws = generate_class_record_sheet(
        wb=wb,
        db=db,
        class_id=class_id,
        subject_id=subject_id,
        academic_period_id=academic_period_id,
        staff_id=staff_id,
    )

    # Fetch scope labels for clean filename
    class_obj = db.get(Class, class_id)
    subj_obj = db.get(Subject, subject_id)

    sec_name = re.sub(r"[^\w\-_.]", "_", (getattr(class_obj, "section_name", "") or "Section").strip())
    sub_name = re.sub(r"[^\w\-_.]", "_", (getattr(subj_obj, "subject_name", "") or "Subject").strip())
    term_name = re.sub(r"[^\w\-_.]", "_", ws.title.strip())

    filename = f"Class_Record_{sec_name}_{sub_name}_{term_name}.xlsx"

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return stream, filename
