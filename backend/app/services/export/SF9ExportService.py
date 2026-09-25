from __future__ import annotations

import io
import re
from calendar import month_abbr
from datetime import date
from typing import Any, cast
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from fastapi import HTTPException
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.academic.AcademicLevel import AcademicLevel
from app.models.academic.AcademicPeriod import AcademicPeriod
from app.models.academic.AcademicYear import AcademicYear
from app.models.academic.Class_ import Class
from app.models.academic.StudentCLass import StudentClass
from app.models.academic.StudentPeriodGrade import StudentPeriodGrade
from app.models.academic.Subject import Subject
from app.models.academic.SubjectLoad import SubjectLoad
from app.models.attendance.Attendance import AttendanceRecord
from app.models.people.AcademicStaff import AcademicStaff
from app.models.people.Student import Student
from app.models.settings.Setting import Setting
from app.services.academic.SubstitutionService import _staff_full_name
from app.services.classes.ClassQueryService import _student_full_name
from app.services.student_record.StudentRecordService import get_performance_descriptor


# ── DepEd Typography & Styling ───────────────────────────────────────

FONT_FAMILY = "Calibri"

FONT_HEADER_GOV = Font(name=FONT_FAMILY, size=9, italic=True, color="333333")
FONT_HEADER_DEPED = Font(name=FONT_FAMILY, size=11, bold=True, color="1F497D")
FONT_TITLE = Font(name=FONT_FAMILY, size=13, bold=True, color="1F497D")
FONT_SECTION_HEADER = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

FONT_META_LABEL = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
FONT_META_VAL = Font(name=FONT_FAMILY, size=9, bold=False, color="000000")

FONT_TABLE_HEAD = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
FONT_CELL_TEXT = Font(name=FONT_FAMILY, size=9, bold=False, color="000000")
FONT_CELL_BOLD = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
FONT_CELL_NUM = Font(name=FONT_FAMILY, size=9, bold=False, color="000000")
FONT_CELL_MUTED = Font(name=FONT_FAMILY, size=8, italic=True, color="555555")

FILL_NAVY = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
FILL_HEADER = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
FILL_SUMMARY = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_MUTED = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

BORDER_THIN = Side(border_style="thin", color="000000")
BORDER_MEDIUM = Side(border_style="medium", color="000000")
BORDER_DOUBLE = Side(border_style="double", color="000000")

CELL_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)
SUMMARY_BORDER = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_DOUBLE)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_NAME = Alignment(horizontal="left", vertical="center", indent=1)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _sanitize_sheet_name(name: str, used_names: set[str]) -> str:
    """Sanitize worksheet name to <= 31 chars and remove illegal Excel chars."""
    cleaned = re.sub(r"[\\/*?:\[\]]", "", name).strip()
    if not cleaned:
        cleaned = "Student"
    candidate = cleaned[:30]
    counter = 1
    while candidate.lower() in used_names:
        suffix = f"_{counter}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def _get_school_settings(db: Session) -> dict[str, str]:
    settings_rows = {str(s.key): str(s.value or "").strip() for s in db.query(Setting).all()}
    return {
        "school_name": settings_rows.get("school_name", "Medellin National Science and Technology School (MNSTS)"),
        "school_id": settings_rows.get("school_id", "303012"),
        "region": settings_rows.get("school_region") or settings_rows.get("region", "Region IV"),
        "division": settings_rows.get("school_division") or settings_rows.get("division", "Fourth District"),
        "district": settings_rows.get("school_district") or settings_rows.get("district", "Medellin"),
        "principal_name": settings_rows.get("principal_name") or settings_rows.get("school_head", "Principal / School Head"),
    }


def _render_sf9_worksheet(
    ws: Worksheet,
    db: Session,
    student: Student,
    class_obj: Class,
    adviser_staff: AcademicStaff | None,
    academic_year: AcademicYear,
    periods: list[AcademicPeriod],
    subjects: list[Subject],
    school_info: dict[str, str],
) -> None:
    ws.views.sheetView[0].showGridLines = True

    # 1. Page Setup
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Calculate columns based on periods count
    num_periods = max(1, len(periods))
    # Layout has max width across sections. Let's establish 8-10 standard columns.
    max_cols = max(7, 2 + num_periods + 2)

    current_row = 1

    # 2. Institutional Header Block
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = "Republic of the Philippines"
    ws.cell(current_row, 1).font = FONT_HEADER_GOV
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = "Department of Education"
    ws.cell(current_row, 1).font = FONT_HEADER_DEPED
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 1

    region_div = f"{school_info['region']} • {school_info['division']}"
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = region_div
    ws.cell(current_row, 1).font = FONT_CELL_BOLD
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = school_info["school_name"].upper()
    ws.cell(current_row, 1).font = FONT_TITLE
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = f"School ID: {school_info['school_id']}"
    ws.cell(current_row, 1).font = FONT_CELL_MUTED
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    ws.cell(current_row, 1).value = "LEARNER'S PROGRESS REPORT CARD (SF9 / Form 138)"
    ws.cell(current_row, 1).font = FONT_TITLE
    ws.cell(current_row, 1).alignment = ALIGN_CENTER
    current_row += 2

    # 3. Learner Profile Metadata
    level_obj = db.get(AcademicLevel, class_obj.academic_level_id) if class_obj.academic_level_id else None
    grade_level_label = f"Grade {level_obj.grade_level}" if level_obj and level_obj.grade_level else (level_obj.level_name if level_obj else "Grade")
    section_display = f"{grade_level_label} - {class_obj.section_name}"

    student_full_name = _student_full_name(student)
    adviser_name = _staff_full_name(adviser_staff) if adviser_staff else "Adviser"

    # Row 1 of metadata: Name & LRN
    ws.cell(current_row, 1).value = "Name:"
    ws.cell(current_row, 1).font = FONT_META_LABEL
    ws.cell(current_row, 2).value = student_full_name
    ws.cell(current_row, 2).font = FONT_META_VAL

    ws.cell(current_row, max_cols - 2).value = "LRN:"
    ws.cell(current_row, max_cols - 2).font = FONT_META_LABEL
    ws.cell(current_row, max_cols - 1).value = student.student_lrn or "N/A"
    ws.cell(current_row, max_cols - 1).font = FONT_META_VAL
    current_row += 1

    # Row 2: Grade & Section, School Year
    ws.cell(current_row, 1).value = "Grade & Section:"
    ws.cell(current_row, 1).font = FONT_META_LABEL
    ws.cell(current_row, 2).value = section_display
    ws.cell(current_row, 2).font = FONT_META_VAL

    ws.cell(current_row, max_cols - 2).value = "School Year:"
    ws.cell(current_row, max_cols - 2).font = FONT_META_LABEL
    ws.cell(current_row, max_cols - 1).value = academic_year.year_label or "SY"
    ws.cell(current_row, max_cols - 1).font = FONT_META_VAL
    current_row += 1

    # Row 3: Sex / Gender, Class Adviser
    ws.cell(current_row, 1).value = "Sex:"
    ws.cell(current_row, 1).font = FONT_META_LABEL
    ws.cell(current_row, 2).value = student.gender or "N/A"
    ws.cell(current_row, 2).font = FONT_META_VAL

    ws.cell(current_row, max_cols - 2).value = "Class Adviser:"
    ws.cell(current_row, max_cols - 2).font = FONT_META_LABEL
    ws.cell(current_row, max_cols - 1).value = adviser_name
    ws.cell(current_row, max_cols - 1).font = FONT_META_VAL
    current_row += 2

    # 4. REPORT ON LEARNING PROGRESS AND ACHIEVEMENT (ACADEMIC GRADES)
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    header_cell = ws.cell(current_row, 1)
    header_cell.value = "REPORT ON LEARNING PROGRESS AND ACHIEVEMENT"
    header_cell.font = FONT_SECTION_HEADER
    header_cell.fill = FILL_NAVY
    header_cell.alignment = ALIGN_CENTER
    current_row += 1

    # Table Header Row 1
    table_head_row = current_row
    ws.cell(table_head_row, 1).value = "LEARNING AREAS"
    ws.cell(table_head_row, 1).font = FONT_TABLE_HEAD
    ws.cell(table_head_row, 1).fill = FILL_HEADER
    ws.cell(table_head_row, 1).border = CELL_BORDER
    ws.cell(table_head_row, 1).alignment = ALIGN_NAME

    # Merge learning areas over cols 1 & 2 for ample width
    learning_area_cols = 2
    ws.merge_cells(start_row=table_head_row, start_column=1, end_row=table_head_row, end_column=learning_area_cols)
    ws.cell(table_head_row, 2).border = CELL_BORDER

    # Period Columns
    col_idx = learning_area_cols + 1
    for p in periods:
        p_cell = ws.cell(table_head_row, col_idx)
        # Period label (e.g. 1, 2, 3, 4 or Q1, Q2)
        p_cell.value = str(p.period_sequence or p.period_name)
        p_cell.font = FONT_TABLE_HEAD
        p_cell.fill = FILL_HEADER
        p_cell.border = CELL_BORDER
        p_cell.alignment = ALIGN_CENTER
        col_idx += 1

    # Final Rating Column
    final_col_idx = col_idx
    ws.cell(table_head_row, final_col_idx).value = "Final Rating"
    ws.cell(table_head_row, final_col_idx).font = FONT_TABLE_HEAD
    ws.cell(table_head_row, final_col_idx).fill = FILL_HEADER
    ws.cell(table_head_row, final_col_idx).border = CELL_BORDER
    ws.cell(table_head_row, final_col_idx).alignment = ALIGN_CENTER
    col_idx += 1

    # Remarks Column
    remarks_col_idx = col_idx
    ws.cell(table_head_row, remarks_col_idx).value = "Remarks"
    ws.cell(table_head_row, remarks_col_idx).font = FONT_TABLE_HEAD
    ws.cell(table_head_row, remarks_col_idx).fill = FILL_HEADER
    ws.cell(table_head_row, remarks_col_idx).border = CELL_BORDER
    ws.cell(table_head_row, remarks_col_idx).alignment = ALIGN_CENTER

    current_row += 1

    # Fetch period grades for this student
    period_ids = [p.academic_period_id for p in periods]
    period_grades = (
        db.query(StudentPeriodGrade)
        .filter(
            StudentPeriodGrade.student_id == student.student_id,
            StudentPeriodGrade.class_id == class_obj.class_id,
            StudentPeriodGrade.academic_period_id.in_(period_ids),
        )
        .all()
    )
    grade_lookup: dict[tuple[int, int], StudentPeriodGrade] = {
        (pg.subject_id, pg.academic_period_id): pg for pg in period_grades
    }

    final_ratings: list[float] = []
    grades_start_row = current_row

    for subj in subjects:
        ws.cell(current_row, 1).value = subj.subject_name
        ws.cell(current_row, 1).font = FONT_CELL_TEXT
        ws.cell(current_row, 1).border = CELL_BORDER
        ws.cell(current_row, 1).alignment = ALIGN_NAME

        ws.cell(current_row, 2).border = CELL_BORDER
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)

        c_idx = 3
        subj_period_values = []
        for p in periods:
            pg = grade_lookup.get((subj.subject_id, p.academic_period_id))
            p_val = float(pg.final_period_grade) if pg and pg.is_finalized and pg.final_period_grade is not None else None
            cell = ws.cell(current_row, c_idx)
            if p_val is not None:
                cell.value = round(p_val, 1)
                cell.font = FONT_CELL_NUM
                subj_period_values.append(p_val)
            else:
                cell.value = ""
                cell.font = FONT_CELL_MUTED
            cell.border = CELL_BORDER
            cell.alignment = ALIGN_CENTER
            c_idx += 1

        # Calculate Final Rating for subject
        final_cell = ws.cell(current_row, final_col_idx)
        final_cell.border = CELL_BORDER
        final_cell.alignment = ALIGN_CENTER

        remarks_cell = ws.cell(current_row, remarks_col_idx)
        remarks_cell.border = CELL_BORDER
        remarks_cell.alignment = ALIGN_CENTER

        if subj_period_values and len(subj_period_values) == len(periods):
            subj_final = round(sum(subj_period_values) / len(subj_period_values), 1)
            final_ratings.append(subj_final)
            final_cell.value = subj_final
            final_cell.font = FONT_CELL_BOLD

            is_passed = subj_final >= 75.0
            remarks_cell.value = "Passed" if is_passed else "Failed"
            remarks_cell.font = FONT_CELL_BOLD if not is_passed else FONT_CELL_TEXT
        elif subj_period_values:
            # Partial terms available
            partial_avg = round(sum(subj_period_values) / len(subj_period_values), 1)
            final_cell.value = partial_avg
            final_cell.font = FONT_CELL_MUTED
            remarks_cell.value = ""
        else:
            final_cell.value = ""
            remarks_cell.value = ""

        current_row += 1

    grades_end_row = current_row - 1

    # General Average Row
    gwa_row = current_row
    ws.cell(gwa_row, 1).value = "General Average (GWA)"
    ws.cell(gwa_row, 1).font = FONT_TABLE_HEAD
    ws.cell(gwa_row, 1).fill = FILL_SUMMARY
    ws.cell(gwa_row, 1).border = SUMMARY_BORDER
    ws.cell(gwa_row, 1).alignment = ALIGN_NAME

    ws.cell(gwa_row, 2).fill = FILL_SUMMARY
    ws.cell(gwa_row, 2).border = SUMMARY_BORDER
    ws.merge_cells(start_row=gwa_row, start_column=1, end_row=gwa_row, end_column=2)

    # Empty fill for term cols
    for c in range(3, final_col_idx):
        ws.cell(gwa_row, c).fill = FILL_SUMMARY
        ws.cell(gwa_row, c).border = SUMMARY_BORDER

    gwa_cell = ws.cell(gwa_row, final_col_idx)
    gwa_cell.fill = FILL_SUMMARY
    gwa_cell.border = SUMMARY_BORDER
    gwa_cell.alignment = ALIGN_CENTER

    gwa_remarks_cell = ws.cell(gwa_row, remarks_col_idx)
    gwa_remarks_cell.fill = FILL_SUMMARY
    gwa_remarks_cell.border = SUMMARY_BORDER
    gwa_remarks_cell.alignment = ALIGN_CENTER

    if final_ratings and len(final_ratings) == len(subjects):
        overall_gwa = round(sum(final_ratings) / len(final_ratings), 1)
        gwa_cell.value = overall_gwa
        gwa_cell.font = FONT_TABLE_HEAD
        gwa_passed = overall_gwa >= 75.0
        gwa_remarks_cell.value = "Passed" if gwa_passed else "Failed"
        gwa_remarks_cell.font = FONT_TABLE_HEAD
    elif final_ratings:
        overall_gwa = round(sum(final_ratings) / len(final_ratings), 1)
        gwa_cell.value = overall_gwa
        gwa_cell.font = FONT_CELL_MUTED
    else:
        gwa_cell.value = ""

    current_row += 1

    # Grading Scale Legend
    legend_start = current_row
    ws.cell(legend_start, 1).value = "Descriptors:"
    ws.cell(legend_start, 1).font = FONT_META_LABEL
    ws.cell(legend_start, 2).value = "Grading Scale:"
    ws.cell(legend_start, 2).font = FONT_META_LABEL
    ws.cell(legend_start, 3).value = "Remarks:"
    ws.cell(legend_start, 3).font = FONT_META_LABEL

    legend_items = [
        ("Advancing", "90 – 100", "Passed"),
        ("Benchmarking", "80 – 89", "Passed"),
        ("Connecting", "75 – 79", "Passed"),
        ("Developing", "65 – 74", "Failed"),
        ("Emerging", "Below 65", "Failed"),
    ]
    leg_row = legend_start + 1
    for desc, scale, rem in legend_items:
        ws.cell(leg_row, 1).value = desc
        ws.cell(leg_row, 1).font = FONT_CELL_TEXT
        ws.cell(leg_row, 2).value = scale
        ws.cell(leg_row, 2).font = FONT_CELL_TEXT
        ws.cell(leg_row, 3).value = rem
        ws.cell(leg_row, 3).font = FONT_CELL_TEXT
        leg_row += 1

    current_row = leg_row + 1

    # 5. REPORT ON LEARNER'S OBSERVED VALUES (CORE VALUES)
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cv_header = ws.cell(current_row, 1)
    cv_header.value = "REPORT ON LEARNER'S OBSERVED VALUES"
    cv_header.font = FONT_SECTION_HEADER
    cv_header.fill = FILL_NAVY
    cv_header.alignment = ALIGN_CENTER
    current_row += 1

    # Table Header
    ws.cell(current_row, 1).value = "Core Values"
    ws.cell(current_row, 1).font = FONT_TABLE_HEAD
    ws.cell(current_row, 1).fill = FILL_HEADER
    ws.cell(current_row, 1).border = CELL_BORDER
    ws.cell(current_row, 1).alignment = ALIGN_LEFT

    ws.cell(current_row, 2).value = "Behavior Statements"
    ws.cell(current_row, 2).font = FONT_TABLE_HEAD
    ws.cell(current_row, 2).fill = FILL_HEADER
    ws.cell(current_row, 2).border = CELL_BORDER
    ws.cell(current_row, 2).alignment = ALIGN_LEFT

    col_idx = 3
    for p in periods:
        p_c = ws.cell(current_row, col_idx)
        p_c.value = str(p.period_sequence or p.period_name)
        p_c.font = FONT_TABLE_HEAD
        p_c.fill = FILL_HEADER
        p_c.border = CELL_BORDER
        p_c.alignment = ALIGN_CENTER
        col_idx += 1

    for c in range(col_idx, max_cols + 1):
        ws.cell(current_row, c).fill = FILL_HEADER
        ws.cell(current_row, c).border = CELL_BORDER

    current_row += 1

    core_values_data = [
        ("Maka-Diyos", "1. Expresses one's spiritual beliefs while respecting the spiritual beliefs of others."),
        ("Maka-Diyos", "2. Shows adherence to ethical principles by upholding truth."),
        ("Makatao", "1. Is sensitive to individual, social, and cultural differences."),
        ("Makatao", "2. Demonstrates solidarity in the school and community."),
        ("Makakalikasan", "1. Cares for the environment and utilizes resources wisely, judiciously, and economically."),
        ("Makabansa", "1. Demonstrates pride in being a Filipino; exercises the rights and responsibilities of a Filipino citizen."),
        ("Makabansa", "2. Demonstrates appropriate behavior in carrying out activities in the school, community, and country."),
    ]

    for cv_title, behavior in core_values_data:
        ws.cell(current_row, 1).value = cv_title
        ws.cell(current_row, 1).font = FONT_CELL_BOLD
        ws.cell(current_row, 1).border = CELL_BORDER
        ws.cell(current_row, 1).alignment = ALIGN_LEFT

        ws.cell(current_row, 2).value = behavior
        ws.cell(current_row, 2).font = FONT_CELL_TEXT
        ws.cell(current_row, 2).border = CELL_BORDER
        ws.cell(current_row, 2).alignment = ALIGN_LEFT

        for c in range(3, max_cols + 1):
            cell = ws.cell(current_row, c)
            cell.value = ""
            cell.border = CELL_BORDER
            cell.alignment = ALIGN_CENTER

        current_row += 1

    # Core Values Legend
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cv_leg = ws.cell(current_row, 1)
    cv_leg.value = "Marking Code:  AO - Always Observed (95-100%)  |  SO - Sometimes Observed (85-94%)  |  RO - Rarely Observed (75-84%)  |  NO - Not Observed (Below 75%)"
    cv_leg.font = FONT_CELL_MUTED
    cv_leg.alignment = ALIGN_CENTER
    current_row += 2

    # 6. REPORT ON ATTENDANCE
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    att_head = ws.cell(current_row, 1)
    att_head.value = "REPORT ON ATTENDANCE"
    att_head.font = FONT_SECTION_HEADER
    att_head.fill = FILL_NAVY
    att_head.alignment = ALIGN_CENTER
    current_row += 1

    # DepEd standard months (August to May or derived from school year)
    month_names = ["Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Total"]
    num_att_cols = len(month_names)

    # Attendance rows: Month header, Days of School, Days Present, Days Absent
    ws.cell(current_row, 1).value = "Month"
    ws.cell(current_row, 1).font = FONT_TABLE_HEAD
    ws.cell(current_row, 1).fill = FILL_HEADER
    ws.cell(current_row, 1).border = CELL_BORDER
    ws.cell(current_row, 1).alignment = ALIGN_LEFT

    # Query attendance records for this student and class
    att_records = (
        db.query(AttendanceRecord)
        .filter(
            AttendanceRecord.student_id == student.student_id,
            AttendanceRecord.class_id == class_obj.class_id,
        )
        .all()
    )

    # Group attendance by month
    att_by_month: dict[int, dict[str, int]] = {}
    for r in att_records:
        m = r.date.month
        if m not in att_by_month:
            att_by_month[m] = {"present": 0, "absent": 0, "late": 0, "total": 0}
        att_by_month[m]["total"] += 1
        if r.status in ("present", "late"):
            att_by_month[m]["present"] += 1
        elif r.status == "absent":
            att_by_month[m]["absent"] += 1

    # Map month names to month numbers (Aug=8, Sep=9, ..., Dec=12, Jan=1, ..., May=5)
    month_mapping = [8, 9, 10, 11, 12, 1, 2, 3, 4, 5]

    for idx, m_label in enumerate(month_names):
        m_cell = ws.cell(current_row, idx + 2)
        m_cell.value = m_label
        m_cell.font = FONT_TABLE_HEAD
        m_cell.fill = FILL_HEADER
        m_cell.border = CELL_BORDER
        m_cell.alignment = ALIGN_CENTER

    current_row += 1

    att_labels = [
        ("No. of school days", "total"),
        ("No. of days present", "present"),
        ("No. of days absent", "absent"),
    ]

    for label, stat_key in att_labels:
        ws.cell(current_row, 1).value = label
        ws.cell(current_row, 1).font = FONT_CELL_TEXT
        ws.cell(current_row, 1).border = CELL_BORDER
        ws.cell(current_row, 1).alignment = ALIGN_LEFT

        row_sum = 0
        for idx, m_num in enumerate(month_mapping):
            cell = ws.cell(current_row, idx + 2)
            cell.border = CELL_BORDER
            cell.alignment = ALIGN_CENTER
            val = att_by_month.get(m_num, {}).get(stat_key, 0)
            if val > 0:
                cell.value = val
                cell.font = FONT_CELL_NUM
                row_sum += val
            else:
                cell.value = ""
                cell.font = FONT_CELL_MUTED

        # Total column
        tot_cell = ws.cell(current_row, len(month_mapping) + 2)
        tot_cell.value = row_sum if row_sum > 0 else ""
        tot_cell.font = FONT_CELL_BOLD
        tot_cell.border = CELL_BORDER
        tot_cell.alignment = ALIGN_CENTER
        tot_cell.fill = FILL_SUBHEADER

        current_row += 1

    current_row += 2

    # 7. CERTIFICATE OF TRANSFER & SIGNATURES BLOCK
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cert_title = ws.cell(current_row, 1)
    cert_title.value = "CERTIFICATE OF TRANSFER"
    cert_title.font = FONT_TABLE_HEAD
    cert_title.alignment = ALIGN_CENTER
    current_row += 1

    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_cols)
    cert_body1 = ws.cell(current_row, 1)
    cert_body1.value = "Eligible for admission to Grade: ____________________     Has advanced credits in: ____________________"
    cert_body1.font = FONT_CELL_TEXT
    cert_body1.alignment = ALIGN_CENTER
    current_row += 2

    # Signatures: Left = Adviser, Right = Principal
    sig_mid = max(3, max_cols // 2)

    ws.cell(current_row, 1).value = adviser_name.upper()
    ws.cell(current_row, 1).font = FONT_CELL_BOLD
    ws.cell(current_row, 1).alignment = ALIGN_CENTER

    ws.cell(current_row, sig_mid + 1).value = school_info["principal_name"].upper()
    ws.cell(current_row, sig_mid + 1).font = FONT_CELL_BOLD
    ws.cell(current_row, sig_mid + 1).alignment = ALIGN_CENTER
    current_row += 1

    ws.cell(current_row, 1).value = "Class Adviser"
    ws.cell(current_row, 1).font = FONT_CELL_MUTED
    ws.cell(current_row, 1).alignment = ALIGN_CENTER

    ws.cell(current_row, sig_mid + 1).value = "Principal / School Head"
    ws.cell(current_row, sig_mid + 1).font = FONT_CELL_MUTED
    ws.cell(current_row, sig_mid + 1).alignment = ALIGN_CENTER

    # 8. Column Width Adjustments
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    for col_letter_idx in range(3, max(max_cols + 1, len(month_names) + 3)):
        col_letter = get_column_letter(col_letter_idx)
        ws.column_dimensions[col_letter].width = 12


def export_student_sf9(
    db: Session,
    class_id: int,
    student_id: uuid.UUID,
    staff_id: str,
) -> io.BytesIO:
    """Generate official SF9 report card workbook for a single student."""
    class_obj = db.get(Class, class_id)
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class section not found.")
    if class_obj.adviser_staff_id != staff_id:
        raise HTTPException(status_code=403, detail="You are not authorized as the adviser of this section.")

    student = db.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    academic_year = db.get(AcademicYear, class_obj.academic_year_id)
    if not academic_year:
        raise HTTPException(status_code=400, detail="Academic year configuration missing.")

    adviser_staff = db.get(AcademicStaff, staff_id)
    school_info = _get_school_settings(db)

    periods = (
        db.query(AcademicPeriod)
        .filter(AcademicPeriod.academic_year_id == class_obj.academic_year_id)
        .order_by(AcademicPeriod.period_sequence.asc(), AcademicPeriod.start_date.asc())
        .all()
    )

    subject_load_rows = (
        db.query(Subject)
        .join(SubjectLoad, Subject.subject_id == SubjectLoad.subject_id)
        .filter(
            SubjectLoad.class_id == class_obj.class_id,
            SubjectLoad.is_active_version.is_(True),
            func.lower(func.coalesce(SubjectLoad.status, "active")).in_(["active", "published"]),
        )
        .order_by(func.lower(Subject.subject_name))
        .all()
    )
    # Deduplicate subjects
    seen = set()
    subjects = []
    for s in subject_load_rows:
        if s.subject_id not in seen:
            seen.add(s.subject_id)
            subjects.append(s)

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    used_names: set[str] = set()
    sheet_title = _sanitize_sheet_name(f"{student.last_name}_{student.first_name[:5]}", used_names)
    ws.title = sheet_title

    _render_sf9_worksheet(
        ws=ws,
        db=db,
        student=student,
        class_obj=class_obj,
        adviser_staff=adviser_staff,
        academic_year=academic_year,
        periods=periods,
        subjects=subjects,
        school_info=school_info,
    )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_advisory_class_sf9(
    db: Session,
    class_id: int,
    staff_id: str,
) -> io.BytesIO:
    """Generate official SF9 report card workbook for the entire advisory class roster."""
    class_obj = db.get(Class, class_id)
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class section not found.")
    if class_obj.adviser_staff_id != staff_id:
        raise HTTPException(status_code=403, detail="You are not authorized as the adviser of this section.")

    academic_year = db.get(AcademicYear, class_obj.academic_year_id)
    if not academic_year:
        raise HTTPException(status_code=400, detail="Academic year configuration missing.")

    adviser_staff = db.get(AcademicStaff, staff_id)
    school_info = _get_school_settings(db)

    periods = (
        db.query(AcademicPeriod)
        .filter(AcademicPeriod.academic_year_id == class_obj.academic_year_id)
        .order_by(AcademicPeriod.period_sequence.asc(), AcademicPeriod.start_date.asc())
        .all()
    )

    subject_load_rows = (
        db.query(Subject)
        .join(SubjectLoad, Subject.subject_id == SubjectLoad.subject_id)
        .filter(
            SubjectLoad.class_id == class_obj.class_id,
            SubjectLoad.is_active_version.is_(True),
            func.lower(func.coalesce(SubjectLoad.status, "active")).in_(["active", "published"]),
        )
        .order_by(func.lower(Subject.subject_name))
        .all()
    )
    seen = set()
    subjects = []
    for s in subject_load_rows:
        if s.subject_id not in seen:
            seen.add(s.subject_id)
            subjects.append(s)

    # Query all enrolled students
    student_rows = (
        db.query(Student)
        .join(StudentClass, Student.student_id == StudentClass.student_id)
        .filter(
            StudentClass.class_id == class_obj.class_id,
            StudentClass.academic_year_id == class_obj.academic_year_id,
            func.lower(func.coalesce(StudentClass.enrollment_status, "enrolled")) == "enrolled",
        )
        .order_by(Student.gender.desc(), Student.last_name.asc(), Student.first_name.asc())
        .all()
    )

    if not student_rows:
        raise HTTPException(status_code=400, detail="No enrolled students found in this advisory class.")

    wb = openpyxl.Workbook()
    used_names: set[str] = set()

    for idx, student in enumerate(student_rows):
        if idx == 0:
            ws = wb.active
            assert ws is not None
        else:
            ws = cast(Worksheet, wb.create_sheet())

        prefix = f"{idx + 1}. "
        sheet_title = _sanitize_sheet_name(f"{prefix}{student.last_name}, {student.first_name[:4]}", used_names)
        ws.title = sheet_title

        _render_sf9_worksheet(
            ws=ws,
            db=db,
            student=student,
            class_obj=class_obj,
            adviser_staff=adviser_staff,
            academic_year=academic_year,
            periods=periods,
            subjects=subjects,
            school_info=school_info,
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def gather_student_sf9_data(
    db: Session,
    class_id: int,
    student_id: str,
    staff_id: str,
) -> dict[str, Any]:
    """Compile structured JSON data for a student's DepEd SF9 report card."""
    class_obj = db.get(Class, class_id)
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class section not found.")
    if class_obj.adviser_staff_id != staff_id:
        raise HTTPException(status_code=403, detail="You are not authorized as the adviser of this section.")

    try:
        student_uuid = uuid.UUID(student_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid student ID format.")

    student = db.get(Student, student_uuid)
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    academic_year = db.get(AcademicYear, class_obj.academic_year_id)
    if not academic_year:
        raise HTTPException(status_code=400, detail="Academic year configuration missing.")

    adviser_staff = db.get(AcademicStaff, staff_id)
    school_info = _get_school_settings(db)

    periods = (
        db.query(AcademicPeriod)
        .filter(AcademicPeriod.academic_year_id == class_obj.academic_year_id)
        .order_by(AcademicPeriod.period_sequence.asc(), AcademicPeriod.start_date.asc())
        .all()
    )

    subject_load_rows = (
        db.query(Subject)
        .join(SubjectLoad, Subject.subject_id == SubjectLoad.subject_id)
        .filter(
            SubjectLoad.class_id == class_obj.class_id,
            SubjectLoad.is_active_version.is_(True),
            func.lower(func.coalesce(SubjectLoad.status, "active")).in_(["active", "published"]),
        )
        .order_by(func.lower(Subject.subject_name))
        .all()
    )
    seen = set()
    subjects = []
    for s in subject_load_rows:
        if s.subject_id not in seen:
            seen.add(s.subject_id)
            subjects.append(s)

    level_obj = db.get(AcademicLevel, class_obj.academic_level_id) if class_obj.academic_level_id else None
    grade_level_label = f"Grade {level_obj.grade_level}" if level_obj and level_obj.grade_level else (level_obj.level_name if level_obj else "Grade")

    # Age calculation
    age: int | None = None
    if student.dob:
        today = date.today()
        age = today.year - student.dob.year - ((today.month, today.day) < (student.dob.month, student.dob.day))

    # Period grades
    period_ids = [p.academic_period_id for p in periods]
    period_grades = (
        db.query(StudentPeriodGrade)
        .filter(
            StudentPeriodGrade.student_id == student.student_id,
            StudentPeriodGrade.class_id == class_obj.class_id,
            StudentPeriodGrade.academic_period_id.in_(period_ids),
        )
        .all()
    )
    grade_lookup = {(pg.subject_id, pg.academic_period_id): pg for pg in period_grades}

    core_subjects = []
    elective_subjects = []
    final_ratings = []

    for subj in subjects:
        subj_grades: dict[str, float | None] = {}
        subj_vals: list[float] = []
        for idx, p in enumerate(periods):
            seq_key = str(p.period_sequence or idx + 1)
            pg = grade_lookup.get((subj.subject_id, p.academic_period_id))
            if pg and pg.is_finalized and pg.final_period_grade is not None:
                val = round(float(pg.final_period_grade), 1)
                subj_grades[seq_key] = val
                subj_vals.append(val)
            else:
                subj_grades[seq_key] = None

        final_rating = None
        remarks = ""
        if subj_vals and len(subj_vals) == len(periods):
            final_rating = round(sum(subj_vals) / len(subj_vals), 1)
            final_ratings.append(final_rating)
            remarks = "Passed" if final_rating >= 75.0 else "Failed"
        elif subj_vals:
            final_rating = round(sum(subj_vals) / len(subj_vals), 1)
            remarks = "Passed" if final_rating >= 75.0 else "Failed"

        item = {
            "subject_id": subj.subject_id,
            "subject_name": subj.subject_name,
            "grades": subj_grades,
            "final_grade": final_rating,
            "remarks": remarks,
        }

        # Check if elective or core
        is_elective = "elective" in subj.subject_name.lower()
        if is_elective:
            elective_subjects.append(item)
        else:
            core_subjects.append(item)

    gwa = None
    gwa_descriptor = None
    gwa_remarks = ""
    if final_ratings:
        gwa = round(sum(final_ratings) / len(final_ratings), 1)
        gwa_descriptor = get_performance_descriptor(gwa)
        gwa_remarks = "Passed" if gwa >= 75.0 else "Failed"

    # Attendance
    att_records = (
        db.query(AttendanceRecord)
        .filter(
            AttendanceRecord.student_id == student.student_id,
            AttendanceRecord.class_id == class_obj.class_id,
        )
        .all()
    )
    att_by_month: dict[int, dict[str, int]] = {}
    for r in att_records:
        m = r.date.month
        if m not in att_by_month:
            att_by_month[m] = {"present": 0, "absent": 0, "total": 0}
        att_by_month[m]["total"] += 1
        if r.status in ("present", "late"):
            att_by_month[m]["present"] += 1
        elif r.status == "absent":
            att_by_month[m]["absent"] += 1

    month_order = [
        ("Jun", 6), ("Jul", 7), ("Aug", 8), ("Sep", 9), ("Oct", 10), ("Nov", 11),
        ("Dec", 12), ("Jan", 1), ("Feb", 2), ("Mar", 3), ("Apr", 4)
    ]
    monthly_attendance = []
    tot_days = 0
    tot_pres = 0
    tot_abs = 0
    for label, m_num in month_order:
        data = att_by_month.get(m_num, {"present": 0, "absent": 0, "total": 0})
        monthly_attendance.append({
            "month": label,
            "class_days": data["total"],
            "present": data["present"],
            "absent": data["absent"],
        })
        tot_days += data["total"]
        tot_pres += data["present"]
        tot_abs += data["absent"]

    return {
        "school_info": {
            "school_name": school_info["school_name"],
            "school_id": school_info["school_id"],
            "region": school_info.get("region", "Region I"),
            "division": school_info.get("division", "Schools Division Office"),
            "district": school_info.get("district", "District"),
            "municipality": school_info.get("municipality", "Municipality/City, Province"),
            "principal_name": school_info.get("principal_name", "School Head"),
        },
        "student": {
            "student_id": str(student.student_id),
            "full_name": _student_full_name(student),
            "student_lrn": student.student_lrn or "",
            "gender": student.gender or "",
            "age": age,
            "dob": student.dob.isoformat() if student.dob else None,
            "track": "",
        },
        "class_info": {
            "class_id": class_obj.class_id,
            "section_name": class_obj.section_name,
            "grade_level": grade_level_label,
            "academic_year": academic_year.year_label or "2026-2027",
            "adviser_name": _staff_full_name(adviser_staff) if adviser_staff else "Adviser",
        },
        "periods": [
            {
                "academic_period_id": p.academic_period_id,
                "period_name": p.period_name,
                "period_sequence": p.period_sequence or (i + 1),
            }
            for i, p in enumerate(periods)
        ],
        "learning_areas": {
            "core": core_subjects,
            "electives": elective_subjects,
        },
        "general_average": {
            "final_rating": gwa,
            "descriptor": gwa_descriptor,
            "remarks": gwa_remarks,
        },
        "descriptors": [
            {"scale": "90-100", "description": "Advancing", "remarks": "Passed"},
            {"scale": "80-89", "description": "Benchmarking", "remarks": "Passed"},
            {"scale": "75-79", "description": "Connecting", "remarks": "Passed"},
            {"scale": "65-74", "description": "Developing", "remarks": "Failed"},
            {"scale": "0-64", "description": "Emerging", "remarks": "Failed"},
        ],
        "attendance": {
            "months": monthly_attendance,
            "total": {
                "class_days": tot_days,
                "present": tot_pres,
                "absent": tot_abs,
            },
        },
    }


def gather_advisory_class_sf9_batch_data(
    db: Session,
    class_id: int,
    staff_id: str,
) -> list[dict[str, Any]]:
    """Compile structured JSON data for all enrolled students in the advisory class."""
    class_obj = db.get(Class, class_id)
    if not class_obj:
        raise HTTPException(status_code=404, detail="Class section not found.")
    if class_obj.adviser_staff_id != staff_id:
        raise HTTPException(status_code=403, detail="You are not authorized as the adviser of this section.")

    student_rows = (
        db.query(Student)
        .join(StudentClass, Student.student_id == StudentClass.student_id)
        .filter(
            StudentClass.class_id == class_obj.class_id,
            StudentClass.academic_year_id == class_obj.academic_year_id,
            func.lower(func.coalesce(StudentClass.enrollment_status, "enrolled")) == "enrolled",
        )
        .order_by(Student.gender.desc(), Student.last_name.asc(), Student.first_name.asc())
        .all()
    )

    batch_data = []
    for s in student_rows:
        batch_data.append(gather_student_sf9_data(db, class_id, str(s.student_id), staff_id))
    return batch_data

