from __future__ import annotations

import io
import os
from typing import Any
from sqlalchemy.orm import Session

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.services.export.SF9ExportService import (
    gather_student_sf9_data,
    gather_advisory_class_sf9_batch_data,
)


COLOR_BLACK = "000000"
COLOR_GRAY = "666666"
COLOR_LIGHT_GRAY = "F2F2F2"
COLOR_HEADER_BG = "E8EEF5"
FONT_NAME = "Calibri"


def _find_logo_path(logo_type: str) -> str | None:
    """Find the best available logo file (PNG, JPG, JPEG) across backend and frontend assets."""
    search_dirs = [
        os.path.join(os.path.dirname(__file__), "assets"),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../../frontend/public/assets")),
    ]
    if logo_type == "deped":
        filenames = [
            "deped_logo.png", "deped-logo.png", "deped_logo.jpg", "deped-logo.jpg",
            "deped_logo.jpeg", "deped-logo.jpeg", "deped.png", "deped.jpg"
        ]
    else:
        filenames = [
            "school_logo.png", "school-logo.png", "school_logo.jpg", "school-logo.jpg",
            "school_logo.jpeg", "school-logo.jpeg", "school.png", "school.jpg", "mnsts.png", "mnsts.jpg"
        ]
    for d in search_dirs:
        if os.path.isdir(d):
            for fn in filenames:
                candidate = os.path.join(d, fn)
                if os.path.isfile(candidate) and os.path.getsize(candidate) > 0:
                    return candidate
    return None


def _prevent_row_split(row: Any) -> None:
    """Ensure a table row does not split across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _set_cell_width(cell: Any, width_in_inches: float) -> None:
    """Set explicit cell width in twips/dxa on tcPr and cell.width so Word never squishes columns."""
    cell.width = Inches(width_in_inches)
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("tcW"):
            tcPr.remove(child)
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(width_in_inches * 1440)}" w:type="dxa"/>')
    tcPr.append(tcW)


def _set_cell_border(cell: Any, **kwargs: dict[str, Any]) -> None:
    """Set custom borders for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f"w:{edge}"
            element = parse_xml(
                f'<{tag} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" '
                f'w:sz="{edge_data.get("sz", 4)}" w:space="0" '
                f'w:color="{edge_data.get("color", "000000")}"/>'
            )
            tcBorders.append(element)
    tcPr.append(tcBorders)


def _set_cell_background(cell: Any, hex_color: str) -> None:
    """Set background shading for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell: Any, top: int = 15, bottom: int = 15, left: int = 30, right: int = 30) -> None:
    """Set cell internal margins in twips (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_cell_text(
    cell: Any,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = 6.5,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    color: str | None = None,
) -> Any:
    """Set text on a cell's first paragraph with strictly zero vertical margin and 1.0 line spacing."""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        if color:
            r = int(color[0:2], 16)
            g = int(color[2:4], 16)
            b = int(color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)
    return p


def _add_p(
    container: Any,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = 7.0,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    space_before: float = 0,
    space_after: float = 1,
    line_spacing: float = 1.0,
    color: str | None = None,
) -> Any:
    """Helper to add a styled paragraph to a cell or document with tight line spacing."""
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        if color:
            r = int(color[0:2], 16)
            g = int(color[2:4], 16)
            b = int(color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)
    return p


def render_student_sf9_docx(
    doc: docx.Document,
    data: dict[str, Any],
    comments: dict[str, str] | None = None,
) -> None:
    """Render a single student's SF9 2-column landscape report card onto exactly ONE page."""
    comments = comments or {}
    school = data["school_info"]
    student = data["student"]
    class_info = data["class_info"]
    periods = data["periods"]
    learning_areas = data["learning_areas"]
    gwa_info = data["general_average"]
    attendance = data["attendance"]
    descriptors = data["descriptors"]

    # 1. Master Table: 1 row, 2 columns (5.15 in each, fits cleanly in Letter and A4)
    master_tbl = doc.add_table(rows=1, cols=2)
    master_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    master_tbl.autofit = False
    _prevent_row_split(master_tbl.rows[0])

    col_w_in = 5.15
    for row in master_tbl.rows:
        _set_cell_width(row.cells[0], col_w_in)
        _set_cell_width(row.cells[1], col_w_in)

    left_cell = master_tbl.cell(0, 0)
    right_cell = master_tbl.cell(0, 1)

    _set_cell_margins(left_cell, top=0, bottom=0, left=15, right=45)
    _set_cell_margins(right_cell, top=0, bottom=0, left=45, right=15)
    _set_cell_border(right_cell, left={"val": "single", "sz": 6, "color": "333333"})

    # Clear default empty paragraph in cells
    left_cell._element.remove(left_cell.paragraphs[0]._element)
    right_cell._element.remove(right_cell.paragraphs[0]._element)

    # ══════════════════════════════════════════════════════════════════════
    # LEFT COLUMN: Learner's Performance Report
    # ══════════════════════════════════════════════════════════════════════

    # ── Header Block with DepEd & School Logos ──
    header_tbl = left_cell.add_table(rows=1, cols=3)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False
    _prevent_row_split(header_tbl.rows[0])

    logo_w_in = 0.70
    mid_w_in = 3.75
    _set_cell_width(header_tbl.rows[0].cells[0], logo_w_in)
    _set_cell_width(header_tbl.rows[0].cells[1], mid_w_in)
    _set_cell_width(header_tbl.rows[0].cells[2], logo_w_in)

    deped_logo_path = _find_logo_path("deped")
    school_logo_path = _find_logo_path("school")

    # Left: DepEd Logo (top-aligned, with space_before so top is slightly below "Republic of the Philippines")
    deped_cell = header_tbl.cell(0, 0)
    deped_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_cell_margins(deped_cell, top=0, bottom=0, left=0, right=10)
    p_deped = deped_cell.paragraphs[0]
    p_deped.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deped.paragraph_format.space_before = Pt(5)
    p_deped.paragraph_format.space_after = Pt(0)
    p_deped.paragraph_format.line_spacing = 1.0
    if deped_logo_path:
        try:
            p_deped.add_run().add_picture(deped_logo_path, width=Inches(0.62), height=Inches(0.62))
        except Exception:
            pass

    # Right: School Logo (top-aligned, matching left logo)
    school_cell = header_tbl.cell(0, 2)
    school_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_cell_margins(school_cell, top=0, bottom=0, left=10, right=0)
    p_school = school_cell.paragraphs[0]
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_school.paragraph_format.space_before = Pt(5)
    p_school.paragraph_format.space_after = Pt(0)
    p_school.paragraph_format.line_spacing = 1.0
    if school_logo_path:
        try:
            p_school.add_run().add_picture(school_logo_path, width=Inches(0.62), height=Inches(0.62))
        except Exception:
            pass

    # Center: Institutional Hierarchy + Automatic School Name (NO "School:" prefix)
    mid_cell = header_tbl.cell(0, 1)
    _set_cell_margins(mid_cell, top=0, bottom=0, left=5, right=5)
    _set_cell_text(mid_cell, "Republic of the Philippines", size=6.0, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_p(mid_cell, "Department of Education", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _add_p(mid_cell, f"{school.get('region', 'Region IV')}", size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _add_p(mid_cell, f"SCHOOLS DIVISION OFFICE OF {school.get('division', '').upper()}", size=6.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _add_p(mid_cell, f"District {school.get('district', '')}", size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _add_p(mid_cell, f"{school.get('municipality', 'Municipality/City, Province')}", size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)

    # Automatic School Name without "School:" label
    _add_p(mid_cell, school["school_name"].upper(), size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    # Report Title
    _add_p(left_cell, "LEARNER'S PERFORMANCE REPORT", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1.5, space_after=0)
    _add_p(left_cell, f"School Year {class_info['academic_year']}", size=7.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    # Student metadata table (borderless, tight, fixed column widths)
    meta_tbl = left_cell.add_table(rows=3, cols=3)
    meta_tbl.autofit = False
    m_col0_in, m_col1_in, m_col2_in = 2.55, 1.10, 1.50
    for r in meta_tbl.rows:
        _prevent_row_split(r)
        _set_cell_width(r.cells[0], m_col0_in)
        _set_cell_width(r.cells[1], m_col1_in)
        _set_cell_width(r.cells[2], m_col2_in)

    # Row 1: Name, Age, Sex
    r0 = meta_tbl.rows[0]
    p = r0.cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Name: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(student["full_name"]).font.size = Pt(6.5)

    p = r0.cells[1].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Age: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(str(student["age"]) if student["age"] is not None else "-").font.size = Pt(6.5)

    p = r0.cells[2].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Sex: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(student["gender"] or "-").font.size = Pt(6.5)

    # Row 2: LRN, Grade, Section
    r1 = meta_tbl.rows[1]
    p = r1.cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("LRN: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(student["student_lrn"] or "-").font.size = Pt(6.5)

    p = r1.cells[1].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Grade: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(class_info["grade_level"]).font.size = Pt(6.5)

    p = r1.cells[2].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Section: ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(class_info["section_name"]).font.size = Pt(6.5)

    # Row 3: Track (SHS only)
    r2 = meta_tbl.rows[2]
    p = r2.cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    p.add_run("Track (SHS only): ").bold = True
    p.runs[0].font.size = Pt(6.5)
    p.add_run(student.get("track") or "General / Core").font.size = Pt(6.5)

    # Dear Parents Note (compact)
    _add_p(left_cell, "Dear Parents,", size=6.5, bold=True, space_before=1, space_after=0.5)
    _add_p(
        left_cell,
        "    This Performance Report shows the ability and progress your child has made in the different learning areas as well as his/her core values. The school welcomes you should you desire to know more about your child's progress.",
        size=6.0,
        space_after=2,
        line_spacing=1.0,
    )

    # LEARNING PROGRESS AND ACHIEVEMENT Table
    _add_p(left_cell, "LEARNING PROGRESS AND ACHIEVEMENT", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=1)

    num_p = max(1, len(periods))
    num_grade_cols = 1 + num_p + 2
    grades_tbl = left_cell.add_table(rows=2, cols=num_grade_cols)
    grades_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    grades_tbl.autofit = False

    la_w = 2.45
    term_w = 0.45
    fg_w = 0.65
    rem_w = 0.70

    for r_idx in range(2):
        _prevent_row_split(grades_tbl.rows[r_idx])
        row = grades_tbl.rows[r_idx]
        _set_cell_width(row.cells[0], la_w)
        for idx in range(num_p):
            _set_cell_width(row.cells[1 + idx], term_w)
        _set_cell_width(row.cells[num_p + 1], fg_w)
        _set_cell_width(row.cells[num_p + 2], rem_w)

    # Table Header Row 0 & 1
    grades_tbl.cell(0, 0).merge(grades_tbl.cell(1, 0))
    _set_cell_width(grades_tbl.cell(0, 0), la_w)
    _set_cell_text(grades_tbl.cell(0, 0), "Learning Areas", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    term_master = grades_tbl.cell(0, 1)
    for idx in range(1, num_p):
        term_master.merge(grades_tbl.cell(0, 1 + idx))
    _set_cell_width(term_master, term_w * num_p)
    _set_cell_text(term_master, "TERM", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for idx, per in enumerate(periods):
        cell = grades_tbl.cell(1, 1 + idx)
        _set_cell_width(cell, term_w)
        _set_cell_text(cell, str(per.get("period_sequence", idx + 1)), bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    grades_tbl.cell(0, num_p + 1).merge(grades_tbl.cell(1, num_p + 1))
    _set_cell_width(grades_tbl.cell(0, num_p + 1), fg_w)
    _set_cell_text(grades_tbl.cell(0, num_p + 1), "Final\nGrade", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    grades_tbl.cell(0, num_p + 2).merge(grades_tbl.cell(1, num_p + 2))
    _set_cell_width(grades_tbl.cell(0, num_p + 2), rem_w)
    _set_cell_text(grades_tbl.cell(0, num_p + 2), "Remarks", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for r_idx in range(2):
        for c in grades_tbl.rows[r_idx].cells:
            _set_cell_border(c, top={"sz": 4}, bottom={"sz": 4}, left={"sz": 4}, right={"sz": 4})
            _set_cell_margins(c, top=8, bottom=8, left=12, right=12)
            _set_cell_background(c, COLOR_HEADER_BG)

    def add_subject_row(subj_item: dict[str, Any], is_section_header: bool = False) -> None:
        row = grades_tbl.add_row()
        _prevent_row_split(row)

        if is_section_header:
            row.cells[0].merge(row.cells[num_grade_cols - 1])
            _set_cell_width(row.cells[0], col_w_in)
            _set_cell_border(row.cells[0], top={"sz": 3}, bottom={"sz": 3}, left={"sz": 3}, right={"sz": 3})
            _set_cell_margins(row.cells[0], top=5, bottom=5, left=12, right=12)
            _set_cell_text(row.cells[0], subj_item["subject_name"], bold=True, size=6.5)
            _set_cell_background(row.cells[0], COLOR_LIGHT_GRAY)
            return

        # Crucial: assign explicit cell widths on every single row!
        _set_cell_width(row.cells[0], la_w)
        for idx in range(num_p):
            _set_cell_width(row.cells[1 + idx], term_w)
        _set_cell_width(row.cells[num_p + 1], fg_w)
        _set_cell_width(row.cells[num_p + 2], rem_w)

        for idx in range(num_grade_cols):
            c = row.cells[idx]
            _set_cell_border(c, top={"sz": 3}, bottom={"sz": 3}, left={"sz": 3}, right={"sz": 3})
            _set_cell_margins(c, top=5, bottom=5, left=10, right=10)

        _set_cell_text(row.cells[0], subj_item["subject_name"], size=6.5)

        for p_idx, per in enumerate(periods):
            seq_key = str(per.get("period_sequence", p_idx + 1))
            val = subj_item["grades"].get(seq_key)
            val_str = f"{val:.0f}" if (val is not None and val == int(val)) else (f"{val:.1f}" if val is not None else "")
            _set_cell_text(row.cells[1 + p_idx], val_str, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        fg_val = subj_item.get("final_grade")
        fg_str = f"{fg_val:.0f}" if (fg_val is not None and fg_val == int(fg_val)) else (f"{fg_val:.1f}" if fg_val is not None else "")
        _set_cell_text(row.cells[num_p + 1], fg_str, bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        _set_cell_text(row.cells[num_p + 2], subj_item.get("remarks") or "", size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Core Subjects
    core_list = learning_areas.get("core") or []
    if core_list:
        add_subject_row({"subject_name": "Core Subjects"}, is_section_header=True)
        for s in core_list:
            add_subject_row(s)

    # Electives
    elective_list = learning_areas.get("electives") or []
    if elective_list:
        add_subject_row({"subject_name": "Elective Subjects"}, is_section_header=True)
        for s in elective_list:
            add_subject_row(s)

    # General Average Row
    gwa_row = grades_tbl.add_row()
    _prevent_row_split(gwa_row)
    for idx in range(num_grade_cols):
        c = gwa_row.cells[idx]
        _set_cell_border(c, top={"sz": 4}, bottom={"sz": 6, "val": "double"}, left={"sz": 4}, right={"sz": 4})
        _set_cell_margins(c, top=6, bottom=6, left=10, right=10)
        _set_cell_background(c, "FFF8DC")

    gwa_label_cell = gwa_row.cells[0]
    for idx in range(1, num_p + 1):
        gwa_label_cell.merge(gwa_row.cells[idx])

    _set_cell_width(gwa_label_cell, la_w + (num_p * term_w))
    _set_cell_width(gwa_row.cells[num_p + 1], fg_w)
    _set_cell_width(gwa_row.cells[num_p + 2], rem_w)

    _set_cell_text(gwa_label_cell, "General Average ", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    gwa_val = gwa_info.get("final_rating")
    gwa_str = f"{gwa_val:.1f}" if gwa_val is not None else ""
    _set_cell_text(gwa_row.cells[num_p + 1], gwa_str, bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(gwa_row.cells[num_p + 2], gwa_info.get("remarks") or "", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # PERFORMANCE DESCRIPTORS Section (compact)
    _add_p(left_cell, "PERFORMANCE DESCRIPTORS", size=7.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1)
    desc_tbl = left_cell.add_table(rows=1 + len(descriptors), cols=3)
    desc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    desc_tbl.autofit = False

    desc_col_ws = [1.50, 2.15, 1.50]
    for row in desc_tbl.rows:
        _prevent_row_split(row)
        for idx in range(3):
            _set_cell_width(row.cells[idx], desc_col_ws[idx])

    for idx, title in enumerate(["Grading Scale", "Description", "Remarks"]):
        c = desc_tbl.cell(0, idx)
        _set_cell_border(c, top={"sz": 4}, bottom={"sz": 4})
        _set_cell_margins(c, top=4, bottom=4, left=10, right=10)
        _set_cell_text(c, title, bold=True, size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER)

    for idx, d in enumerate(descriptors):
        row = desc_tbl.rows[idx + 1]
        for c_idx, val in enumerate([d["scale"], d["description"], d["remarks"]]):
            c = row.cells[c_idx]
            _set_cell_margins(c, top=3, bottom=3, left=10, right=10)
            _set_cell_text(c, val, size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ══════════════════════════════════════════════════════════════════════
    # RIGHT COLUMN: Attendance, Remarks, Signatures, Transfer
    # ══════════════════════════════════════════════════════════════════════

    # 1. ATTENDANCE RECORD
    _add_p(right_cell, "ATTENDANCE RECORD", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)

    att_months = attendance["months"]
    num_att_cols = 1 + len(att_months) + 1
    att_tbl = right_cell.add_table(rows=4, cols=num_att_cols)
    att_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    att_tbl.autofit = False

    lbl_w = 1.40
    m_w = 0.30
    tot_w = 0.45

    for row in att_tbl.rows:
        _prevent_row_split(row)
        _set_cell_width(row.cells[0], lbl_w)
        for idx in range(len(att_months)):
            _set_cell_width(row.cells[1 + idx], m_w)
        _set_cell_width(row.cells[num_att_cols - 1], tot_w)

    _set_cell_text(att_tbl.cell(0, 0), "Month", bold=True, size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for idx, m in enumerate(att_months):
        _set_cell_text(att_tbl.cell(0, 1 + idx), m["month"], bold=True, size=5.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(att_tbl.cell(0, num_att_cols - 1), "Total", bold=True, size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 1: No. of Class Days
    _set_cell_text(att_tbl.cell(1, 0), "No. of Class Days", size=5.5)
    for idx, m in enumerate(att_months):
        _set_cell_text(att_tbl.cell(1, 1 + idx), str(m["class_days"]) if m["class_days"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(att_tbl.cell(1, num_att_cols - 1), str(attendance["total"]["class_days"]) if attendance["total"]["class_days"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 2: No. of Days Present
    _set_cell_text(att_tbl.cell(2, 0), "No. of Days Present", size=5.5)
    for idx, m in enumerate(att_months):
        _set_cell_text(att_tbl.cell(2, 1 + idx), str(m["present"]) if m["present"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(att_tbl.cell(2, num_att_cols - 1), str(attendance["total"]["present"]) if attendance["total"]["present"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 3: No. of Days Absent
    _set_cell_text(att_tbl.cell(3, 0), "No. of Days Absent", size=5.5)
    for idx, m in enumerate(att_months):
        _set_cell_text(att_tbl.cell(3, 1 + idx), str(m["absent"]) if m["absent"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(att_tbl.cell(3, num_att_cols - 1), str(attendance["total"]["absent"]) if attendance["total"]["absent"] else "", size=5.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for r in att_tbl.rows:
        for c in r.cells:
            _set_cell_border(c, top={"sz": 3}, bottom={"sz": 3}, left={"sz": 3}, right={"sz": 3})
            _set_cell_margins(c, top=6, bottom=6, left=6, right=6)

    # 2. TEACHER'S COMMENTS/REMARKS (compact boxes)
    _add_p(right_cell, "TEACHER'S COMMENTS/REMARKS", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1)

    for term_num, term_key in enumerate(["term1", "term2", "term3"], 1):
        c_tbl = right_cell.add_table(rows=1, cols=1)
        c_tbl.autofit = False
        _prevent_row_split(c_tbl.rows[0])
        _set_cell_width(c_tbl.cell(0, 0), col_w_in)
        c = c_tbl.cell(0, 0)
        _set_cell_border(c, top={"sz": 3}, bottom={"sz": 3}, left={"sz": 3}, right={"sz": 3})
        _set_cell_margins(c, top=8, bottom=10, left=20, right=20)

        t_text = comments.get(term_key) or ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if t_text else 6)
        p.paragraph_format.line_spacing = 1.0
        r_title = p.add_run(f"Term {term_num}: ")
        r_title.bold = True
        r_title.font.name = FONT_NAME
        r_title.font.size = Pt(6.5)
        if t_text:
            r_body = p.add_run(t_text)
            r_body.font.name = FONT_NAME
            r_body.font.size = Pt(6.0)

    # 3. PARENTS/GUARDIAN'S SIGNATURE
    _add_p(right_cell, "PARENTS/GUARDIAN'S SIGNATURE", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1)
    _add_p(right_cell, "Term 1  ___________________________________________________________", size=6.5, space_after=1)
    _add_p(right_cell, "Term 2  ___________________________________________________________", size=6.5, space_after=1)
    _add_p(right_cell, "Term 3  ___________________________________________________________", size=6.5, space_after=2)

    # 4. CERTIFICATE OF TRANSFER
    _add_p(right_cell, "CERTIFICATE OF TRANSFER", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1.5, space_after=0.5)
    _add_p(
        right_cell,
        "This is to certify that the above-named learner has satisfactorily completed the requirements for the grade level indicated.",
        size=6.0,
        space_after=1,
    )
    _add_p(right_cell, "Admitted to Grade: ________________________________________________", size=6.0, space_after=1)
    _add_p(right_cell, "Eligible for Admission to Grade: ___________________________________", size=6.0, space_after=1.5)

    _add_p(right_cell, "Approved:", size=6.5, bold=True, space_after=0.5)
    sig_tbl = right_cell.add_table(rows=2, cols=2)
    sig_tbl.autofit = False
    _prevent_row_split(sig_tbl.rows[0])
    _prevent_row_split(sig_tbl.rows[1])
    _set_cell_width(sig_tbl.columns[0].cells[0], 2.57)
    _set_cell_width(sig_tbl.columns[0].cells[1], 2.57)
    _set_cell_width(sig_tbl.columns[1].cells[0], 2.58)
    _set_cell_width(sig_tbl.columns[1].cells[1], 2.58)

    _set_cell_text(sig_tbl.cell(0, 0), school.get("principal_name") or "School Head", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(sig_tbl.cell(0, 1), class_info.get("adviser_name") or "Adviser", bold=True, size=6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(sig_tbl.cell(1, 0), "_____________________________\nSchool Head", size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(sig_tbl.cell(1, 1), "_____________________________\nAdviser", size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 5. CANCELLATION OF ELIGIBILITY TO TRANSFER
    _add_p(right_cell, "CANCELLATION OF ELIGIBILITY TO TRANSFER", size=7.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=0.5)
    _add_p(right_cell, "Admitted in: _________________________________     Date: _________________", size=6.0, space_after=1)
    _add_p(right_cell, "_____________________________\nSchool Head", size=6.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)


def export_student_sf9_docx(
    db: Session,
    class_id: int,
    student_id: str,
    staff_id: str,
    comments: dict[str, str] | None = None,
) -> io.BytesIO:
    """Generate Word (.docx) SF9 report card for an individual student fitting strictly on 1 page."""
    data = gather_student_sf9_data(db, class_id, student_id, staff_id)

    doc = docx.Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.2)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.25)
    section.right_margin = Inches(0.25)

    render_student_sf9_docx(doc, data, comments)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def export_advisory_class_sf9_batch_docx(
    db: Session,
    class_id: int,
    staff_id: str,
) -> io.BytesIO:
    """Generate multi-student Word (.docx) workbook with a 2-column landscape page per student."""
    batch_data = gather_advisory_class_sf9_batch_data(db, class_id, staff_id)

    doc = docx.Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.2)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.25)
    section.right_margin = Inches(0.25)

    for idx, student_data in enumerate(batch_data):
        if idx > 0:
            doc.add_page_break()
        render_student_sf9_docx(doc, student_data)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
